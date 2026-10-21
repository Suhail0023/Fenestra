#!/usr/bin/env python3
"""
AZAL_PROJECTS_PROGRESS.xlsx  ->  Primavera P6 XER   (v3 - resource loaded, built on the layout of a real P6 17.7 export)

Usage:  python azal_to_xer.py [input.xlsx] [output_folder]

Output: ONE P6 project ("AZAL") with a WBS per building (Excel PROJECT column),
        each split into SUPPLY / INSTALLATION, 30 activities per building.
Rules : projects whose PROJECT cell is filled pure red are skipped.
Update: re-run after updating the Excel, then File > Import > XER > "Update Existing Project".
Needs : pip install openpyxl   (save the Excel once in Excel so formula results are cached)
"""
import sys, os, re, math, base64, hashlib, datetime as dt
from collections import OrderedDict, Counter
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "/mnt/user-data/uploads/AZAL_PROJECTS_PROGRESS.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "/mnt/user-data/outputs"

# ----------------------------------------------------------------- settings
DATA_DATE = dt.datetime(2026, 10, 16, 0, 0)          # P6 data date
P6_VERSION = "17.7"
PROJECT_ID, PROJECT_NAME = "AZAL", "AZAL PROJECTS"
EPS_PARENT_ID = "3667"                               # same EPS id as in the template export
HRS_PER_DAY = 8
DAY_START, DAY_END = 8, 16                           # calendar 08:00-16:00
NO_DATE_PLACEHOLDER_DAYS = 1

# fixed Activity-ID map so IDs never shift when rows are re-ordered
SUPPLY = ["CURTAIN WALL FRAME PRODUCTION", "CURTAIN WALL GLASS SUPPLY",
          "DOORS AND WINDOWS FRAME PRODUCTION", "DOORS AND WINDOWS SHUTTER PRODUCTION",
          "BALUSTRADE CHANNEL PRODUCTION", "BALUSTRADE GLASS SUPPLY",
          "BALCONY PARTITION CHANNEL PRODUCTION", "BALCONY PARTITION GLASS SUPPLY",
          "SHOWER PARTITION CHANNEL PRODUCTION", "SHOWER PARTITION GLASS SUPPLY",
          "LOUVERS PRODUCTION", "ACP PRODUCTION", "PERGOLA PRODUCTION",
          "DECORATIVE SCREEN PRODUCTION", "STAIRCASE HANDRAIL PRODUCTION"]
INSTALL = ["CURTAIN WALL FRAME INSTALLATION", "CURTAIN WALL GLASS INSTALLATION",
           "DOORS AND WINDOWS FRAME INSTALLATION", "DOORS AND WINDOWS SHUTTER INSTALLATION",
           "BALUSTRADE CHANNEL INSTALLATION", "BALUSTRADE GLASS INSTALLATION",
           "BALCONY PARTITION CHANNEL INSTALLATION", "BALCONY PARTITION GLASS INSTALLATION",
           "SHOWER PARTITION CHANNEL INSTALLATION", "SHOWER PARTITION GLASS INSTALLATION",
           "LOUVERS INSTALLATION", "ACP INSTALLATION", "PERGOLA INSTALLATION",
           "DECORATIVE SCREEN INSTALLATION", "STAIRCASE HANDRAIL INSTALLATION"]
ID_MAP = {d: f"SUP-{i+1:02d}" for i, d in enumerate(SUPPLY)}
ID_MAP.update({d: f"INS-{i+1:02d}" for i, d in enumerate(INSTALL)})

# ---- layout copied from a real P6 17.7 export (AMF_260710.xer) ----
FIELDS = {'CURRTYPE': ['curr_id', 'decimal_digit_cnt', 'curr_symbol', 'decimal_symbol', 'digit_group_symbol', 'pos_curr_fmt_type', 'neg_curr_fmt_type', 'curr_type', 'curr_short_name', 'group_digit_cnt', 'base_exch_rate'], 'OBS': ['obs_id', 'parent_obs_id', 'guid', 'seq_num', 'obs_name', 'obs_descr'], 'PROJECT': ['proj_id', 'fy_start_month_num', 'rsrc_self_add_flag', 'allow_complete_flag', 'rsrc_multi_assign_flag', 'checkout_flag', 'project_flag', 'step_complete_flag', 'cost_qty_recalc_flag', 'batch_sum_flag', 'name_sep_char', 'def_complete_pct_type', 'proj_short_name', 'acct_id', 'orig_proj_id', 'source_proj_id', 'base_type_id', 'clndr_id', 'sum_base_proj_id', 'task_code_base', 'task_code_step', 'priority_num', 'wbs_max_sum_level', 'strgy_priority_num', 'last_checksum', 'critical_drtn_hr_cnt', 'def_cost_per_qty', 'last_recalc_date', 'plan_start_date', 'plan_end_date', 'scd_end_date', 'add_date', 'last_tasksum_date', 'fcst_start_date', 'def_duration_type', 'task_code_prefix', 'guid', 'def_qty_type', 'add_by_name', 'web_local_root_path', 'proj_url', 'def_rate_type', 'add_act_remain_flag', 'act_this_per_link_flag', 'def_task_type', 'act_pct_link_flag', 'critical_path_type', 'task_code_prefix_flag', 'def_rollup_dates_flag', 'use_project_baseline_flag', 'rem_target_link_flag', 'reset_planned_flag', 'allow_neg_act_flag', 'sum_assign_level', 'last_fin_dates_id', 'last_baseline_update_date', 'cr_external_key', 'apply_actuals_date', 'location_id', 'loaded_scope_level', 'export_flag', 'new_fin_dates_id', 'baselines_to_export', 'baseline_names_to_export', 'next_data_date', 'close_period_flag', 'sum_refresh_date', 'trsrcsum_loaded'], 'CALENDAR': ['clndr_id', 'default_flag', 'clndr_name', 'proj_id', 'base_clndr_id', 'last_chng_date', 'clndr_type', 'day_hr_cnt', 'week_hr_cnt', 'month_hr_cnt', 'year_hr_cnt', 'rsrc_private', 'clndr_data'], 'PROJWBS': ['wbs_id', 'proj_id', 'obs_id', 'seq_num', 'est_wt', 'proj_node_flag', 'sum_data_flag', 'status_code', 'wbs_short_name', 'wbs_name', 'phase_id', 'parent_wbs_id', 'ev_user_pct', 'ev_etc_user_value', 'orig_cost', 'indep_remain_total_cost', 'ann_dscnt_rate_pct', 'dscnt_period_type', 'indep_remain_work_qty', 'anticip_start_date', 'anticip_end_date', 'ev_compute_type', 'ev_etc_compute_type', 'guid', 'tmpl_guid', 'plan_open_state'], 'TASK': ['task_id', 'proj_id', 'wbs_id', 'clndr_id', 'phys_complete_pct', 'rev_fdbk_flag', 'est_wt', 'lock_plan_flag', 'auto_compute_act_flag', 'complete_pct_type', 'task_type', 'duration_type', 'status_code', 'task_code', 'task_name', 'rsrc_id', 'total_float_hr_cnt', 'free_float_hr_cnt', 'remain_drtn_hr_cnt', 'act_work_qty', 'remain_work_qty', 'target_work_qty', 'target_drtn_hr_cnt', 'target_equip_qty', 'act_equip_qty', 'remain_equip_qty', 'cstr_date', 'act_start_date', 'act_end_date', 'late_start_date', 'late_end_date', 'expect_end_date', 'early_start_date', 'early_end_date', 'restart_date', 'reend_date', 'target_start_date', 'target_end_date', 'rem_late_start_date', 'rem_late_end_date', 'cstr_type', 'priority_type', 'suspend_date', 'resume_date', 'float_path', 'float_path_order', 'guid', 'tmpl_guid', 'cstr_date2', 'cstr_type2', 'driving_path_flag', 'act_this_per_work_qty', 'act_this_per_equip_qty', 'external_early_start_date', 'external_late_end_date', 'create_date', 'update_date', 'create_user', 'update_user', 'location_id'], 'UMEASURE': ['unit_id', 'seq_num', 'unit_abbrev', 'unit_name'], 'RSRC': ['rsrc_id', 'parent_rsrc_id', 'clndr_id', 'role_id', 'shift_id', 'user_id', 'pobs_id', 'guid', 'rsrc_seq_num', 'email_addr', 'employee_code', 'office_phone', 'other_phone', 'rsrc_name', 'rsrc_short_name', 'rsrc_title_name', 'def_qty_per_hr', 'cost_qty_type', 'ot_factor', 'active_flag', 'auto_compute_act_flag', 'def_cost_qty_link_flag', 'ot_flag', 'curr_id', 'unit_id', 'rsrc_type', 'location_id', 'rsrc_notes', 'load_tasks_flag', 'level_flag', 'last_checksum'], 'RSRCRATE': ['rsrc_rate_id', 'rsrc_id', 'max_qty_per_hr', 'cost_per_qty', 'start_date', 'shift_period_id', 'cost_per_qty2', 'cost_per_qty3', 'cost_per_qty4', 'cost_per_qty5'], 'TASKRSRC': ['taskrsrc_id', 'task_id', 'proj_id', 'cost_qty_link_flag', 'role_id', 'acct_id', 'rsrc_id', 'pobs_id', 'skill_level', 'remain_qty', 'target_qty', 'remain_qty_per_hr', 'target_lag_drtn_hr_cnt', 'target_qty_per_hr', 'act_ot_qty', 'act_reg_qty', 'relag_drtn_hr_cnt', 'ot_factor', 'cost_per_qty', 'target_cost', 'act_reg_cost', 'act_ot_cost', 'remain_cost', 'act_start_date', 'act_end_date', 'restart_date', 'reend_date', 'target_start_date', 'target_end_date', 'rem_late_start_date', 'rem_late_end_date', 'rollup_dates_flag', 'target_crv', 'remain_crv', 'actual_crv', 'ts_pend_act_end_flag', 'guid', 'rate_type', 'act_this_per_cost', 'act_this_per_qty', 'curv_id', 'rsrc_type', 'cost_per_qty_source_type', 'create_user', 'create_date', 'has_rsrchours', 'taskrsrc_sum_id']}
CAL_ROW = {'clndr_id': '989', 'default_flag': 'N', 'clndr_name': 'Corporate - Standard Full Time', 'proj_id': '', 'base_clndr_id': '', 'last_chng_date': '2026-09-25 00:00', 'clndr_type': 'CA_Base', 'day_hr_cnt': '8', 'week_hr_cnt': '40', 'month_hr_cnt': '172', 'year_hr_cnt': '2000', 'rsrc_private': 'N', 'clndr_data': '(0||CalendarData()(\x7f\x7f  (0||DaysOfWeek()(\x7f\x7f    (0||1()())\x7f\x7f    (0||2()(\x7f\x7f      (0||0(s|08:00|f|16:00)())))\x7f\x7f    (0||3()(\x7f\x7f      (0||0(s|08:00|f|16:00)())))\x7f\x7f    (0||4()(\x7f\x7f      (0||0(s|08:00|f|16:00)())))\x7f\x7f    (0||5()(\x7f\x7f      (0||0(s|08:00|f|16:00)())))\x7f\x7f    (0||6()(\x7f\x7f      (0||0(s|08:00|f|16:00)())))\x7f\x7f    (0||7()(\x7f\x7f      (0||0(s|08:00|f|16:00)())))))\x7f\x7f  (0||VIEW(ShowTotal|Y)())\x7f\x7f  (0||Exceptions()(\x7f\x7f    (0||0(d|40179)())\x7f\x7f    (0||1(d|40329)())\x7f\x7f    (0||2(d|40364)())\x7f\x7f    (0||3(d|40427)())\x7f\x7f    (0||4(d|40507)())\x7f\x7f    (0||5(d|40536)())\x7f\x7f    (0||6(d|40543)())\x7f\x7f    (0||7(d|40693)())\x7f\x7f    (0||8(d|40728)())\x7f\x7f    (0||9(d|40791)())\x7f\x7f    (0||10(d|40871)())\x7f\x7f    (0||11(d|40903)())\x7f\x7f    (0||12(d|40910)())\x7f\x7f    (0||13(d|41057)())\x7f\x7f    (0||14(d|41094)())\x7f\x7f    (0||15(d|41155)())\x7f\x7f    (0||16(d|41235)())\x7f\x7f    (0||17(d|41268)())\x7f\x7f    (0||18(d|41275)())\x7f\x7f    (0||19(d|41421)())\x7f\x7f    (0||20(d|41459)())\x7f\x7f    (0||21(d|41519)())\x7f\x7f    (0||22(d|41606)())\x7f\x7f    (0||23(d|41633)())\x7f\x7f    (0||24(d|45904)())\x7f\x7f    (0||25(d|46358)())\x7f\x7f    (0||26(d|46359)())\x7f\x7f    (0||27(d|46388)())\x7f\x7f    (0||28(d|46456)())\x7f\x7f    (0||29(d|46457)())\x7f\x7f    (0||30(d|46458)())\x7f\x7f    (0||31(d|46522)())\x7f\x7f    (0||32(d|46524)())\x7f\x7f    (0||33(d|46525)())\x7f\x7f    (0||34(d|46615)())\x7f\x7f    (0||35(d|46723)())\x7f\x7f    (0||36(d|46724)())\x7f\x7f    (0||37(d|46753)())\x7f\x7f    (0||38(d|46809)())\x7f\x7f    (0||39(d|46811)())\x7f\x7f    (0||40(d|46812)())\x7f\x7f    (0||41(d|46877)())\x7f\x7f    (0||42(d|46878)())\x7f\x7f    (0||43(d|46879)())\x7f\x7f    (0||44(d|46898)())\x7f\x7f    (0||45(d|46968)())\x7f\x7f    (0||46(d|47088)())\x7f\x7f    (0||47(d|47089)())))))'}
PROJECT_DEFAULTS = {'proj_id': '865', 'fy_start_month_num': '1', 'rsrc_self_add_flag': 'Y', 'allow_complete_flag': 'Y', 'rsrc_multi_assign_flag': 'Y', 'checkout_flag': 'N', 'project_flag': 'Y', 'step_complete_flag': 'N', 'cost_qty_recalc_flag': 'N', 'batch_sum_flag': 'Y', 'name_sep_char': '.', 'def_complete_pct_type': 'CP_Phys', 'proj_short_name': 'AM', 'acct_id': '', 'orig_proj_id': '', 'source_proj_id': '', 'base_type_id': '', 'clndr_id': '989', 'sum_base_proj_id': '', 'task_code_base': '1000', 'task_code_step': '10', 'priority_num': '10', 'wbs_max_sum_level': '2', 'strgy_priority_num': '500', 'last_checksum': '', 'critical_drtn_hr_cnt': '0', 'def_cost_per_qty': '0.0000', 'last_recalc_date': '2026-07-10 00:00', 'plan_start_date': '2026-07-10 00:00', 'plan_end_date': '', 'scd_end_date': '2026-07-15 16:00', 'add_date': '2026-10-16 10:30', 'last_tasksum_date': '', 'fcst_start_date': '', 'def_duration_type': 'DT_FixedDUR2', 'task_code_prefix': 'A', 'guid': 'QDtX4Bk7dE2/SwHnf+rvBg', 'def_qty_type': 'QT_Hour', 'add_by_name': 'admin', 'web_local_root_path': '', 'proj_url': '', 'def_rate_type': 'COST_PER_QTY', 'add_act_remain_flag': 'N', 'act_this_per_link_flag': 'Y', 'def_task_type': 'TT_Task', 'act_pct_link_flag': 'Y', 'critical_path_type': 'CT_TotFloat', 'task_code_prefix_flag': 'Y', 'def_rollup_dates_flag': 'Y', 'use_project_baseline_flag': 'Y', 'rem_target_link_flag': 'Y', 'reset_planned_flag': 'N', 'allow_neg_act_flag': 'N', 'sum_assign_level': 'SL_Taskrsrc', 'last_fin_dates_id': '', 'last_baseline_update_date': '', 'cr_external_key': '', 'apply_actuals_date': '', 'location_id': '', 'loaded_scope_level': '7', 'export_flag': 'Y', 'new_fin_dates_id': '', 'baselines_to_export': '', 'baseline_names_to_export': '', 'next_data_date': '', 'close_period_flag': '', 'sum_refresh_date': '', 'trsrcsum_loaded': ''}
USD_ROW = {'curr_id': '1', 'decimal_digit_cnt': '2', 'curr_symbol': '$', 'decimal_symbol': '.', 'digit_group_symbol': ',', 'pos_curr_fmt_type': '#1.1', 'neg_curr_fmt_type': '(#1.1)', 'curr_type': 'US Dollar', 'curr_short_name': 'USD', 'group_digit_cnt': '3', 'base_exch_rate': '1'}
OBS_ROW = {'obs_id': '565', 'parent_obs_id': '', 'guid': '', 'seq_num': '0', 'obs_name': 'Enterprise', 'obs_descr': ''}
UMEASURE_ROWS = [{'unit_id': '100', 'seq_num': '0', 'unit_abbrev': 'SQM', 'unit_name': 'Square meter'}]
LABOR_RSRC_ROWS = [{'rsrc_id': '1375', 'parent_rsrc_id': '', 'clndr_id': '989', 'role_id': '', 'shift_id': '', 'user_id': '', 'pobs_id': '', 'guid': 'eqB40cCzS0OMnbLKnGxPMw', 'rsrc_seq_num': '109', 'email_addr': '', 'employee_code': '', 'office_phone': '', 'other_phone': '', 'rsrc_name': 'PRODUCTION MANHOUR', 'rsrc_short_name': 'PMHR', 'rsrc_title_name': '', 'def_qty_per_hr': '1', 'cost_qty_type': 'QT_Hour', 'ot_factor': '', 'active_flag': 'Y', 'auto_compute_act_flag': 'Y', 'def_cost_qty_link_flag': 'Y', 'ot_flag': 'N', 'curr_id': '1', 'unit_id': '', 'rsrc_type': 'RT_Labor', 'location_id': '', 'rsrc_notes': '', 'load_tasks_flag': '', 'level_flag': '', 'last_checksum': ''}, {'rsrc_id': '1376', 'parent_rsrc_id': '', 'clndr_id': '989', 'role_id': '', 'shift_id': '', 'user_id': '', 'pobs_id': '', 'guid': 'ZcaUfBp2rEqHw5RV76xnBQ', 'rsrc_seq_num': '110', 'email_addr': '', 'employee_code': '', 'office_phone': '', 'other_phone': '', 'rsrc_name': 'INSTALLATION MANHOUR', 'rsrc_short_name': 'INMHR', 'rsrc_title_name': '', 'def_qty_per_hr': '1', 'cost_qty_type': 'QT_Hour', 'ot_factor': '', 'active_flag': 'Y', 'auto_compute_act_flag': 'Y', 'def_cost_qty_link_flag': 'Y', 'ot_flag': 'N', 'curr_id': '1', 'unit_id': '', 'rsrc_type': 'RT_Labor', 'location_id': '', 'rsrc_notes': '', 'load_tasks_flag': '', 'level_flag': '', 'last_checksum': ''}]
LABOR_RATE_ROWS = [{'rsrc_rate_id': '1820', 'rsrc_id': '1375', 'max_qty_per_hr': '60', 'cost_per_qty': '10.0000', 'start_date': '2026-01-01 00:00', 'shift_period_id': '', 'cost_per_qty2': '0.0000', 'cost_per_qty3': '0.0000', 'cost_per_qty4': '0.0000', 'cost_per_qty5': '0.0000'}, {'rsrc_rate_id': '1821', 'rsrc_id': '1376', 'max_qty_per_hr': '32.5', 'cost_per_qty': '1.0000', 'start_date': '2026-01-01 00:00', 'shift_period_id': '', 'cost_per_qty2': '0.0000', 'cost_per_qty3': '0.0000', 'cost_per_qty4': '0.0000', 'cost_per_qty5': '0.0000'}]
MAT_RSRC_SAMPLE = {'rsrc_id': '1370', 'parent_rsrc_id': '', 'clndr_id': '989', 'role_id': '', 'shift_id': '', 'user_id': '', 'pobs_id': '', 'guid': 'Y+A4Gin8i0KFLWVWerlOLg', 'rsrc_seq_num': '101', 'email_addr': '', 'employee_code': '', 'office_phone': '', 'other_phone': '', 'rsrc_name': 'KITCHEN COUNTER TOP SQM', 'rsrc_short_name': 'KCT.SQM', 'rsrc_title_name': '', 'def_qty_per_hr': '1', 'cost_qty_type': 'QT_Hour', 'ot_factor': '0', 'active_flag': 'Y', 'auto_compute_act_flag': 'N', 'def_cost_qty_link_flag': 'Y', 'ot_flag': 'N', 'curr_id': '1', 'unit_id': '100', 'rsrc_type': 'RT_Mat', 'location_id': '', 'rsrc_notes': '', 'load_tasks_flag': '', 'level_flag': '', 'last_checksum': ''}


# ---- relationships (applied inside every building, for both SUPPLY and INSTALLATION activities)
#   (predecessor, successor, type, lag in working days)
LINKS = [("01", "02", "SS", 15),    # CURTAIN WALL FRAME      -> CURTAIN WALL GLASS
         ("05", "06", "SS", 15),    # BALUSTRADE CHANNEL      -> BALUSTRADE GLASS
         ("03", "04", "SS", 15)]    # DOORS & WINDOWS FRAME   -> DOORS & WINDOWS SHUTTER
#   every SUPPLY/PRODUCTION activity -> its INSTALLATION activity (same scope item), SS + 15 working days
CROSS_LINKS = [(f"SUP-{n:02d}", f"INS-{n:02d}", "SS", 15) for n in range(1, 16)]
FIELDS["TASKPRED"] = ["task_pred_id", "task_id", "pred_task_id", "proj_id", "pred_proj_id", "pred_type",
                      "lag_hr_cnt", "comments", "float_path", "aoa_flag"]

# short resource IDs for the quantity (material) resources, one per scope item
ITEM_SHORT = {"CURTAIN WALL FRAME": "CW.FRAME", "CURTAIN WALL GLASS": "CW.GLASS", "DOORS AND WINDOWS FRAME": "DW.FRAME",
              "DOORS AND WINDOWS SHUTTER": "DW.SHUTTER", "BALUSTRADE CHANNEL": "BAL.CHANNEL", "BALUSTRADE GLASS": "BAL.GLASS",
              "BALCONY PARTITION CHANNEL": "BCP.CHANNEL", "BALCONY PARTITION GLASS": "BCP.GLASS",
              "SHOWER PARTITION CHANNEL": "SHP.CHANNEL", "SHOWER PARTITION GLASS": "SHP.GLASS", "LOUVERS": "LOUVERS",
              "ACP": "ACP", "PERGOLA": "PERGOLA", "DECORATIVE SCREEN": "DEC.SCREEN", "STAIRCASE HANDRAIL": "ST.HANDRAIL"}
LABOR = {r["rsrc_short_name"]: r for r in LABOR_RSRC_ROWS}
LABOR_RATE = {r["rsrc_id"]: float(r["cost_per_qty"] or 0) for r in LABOR_RATE_ROWS}
mat_rsrc = OrderedDict()      # "item|uom" -> rsrc record
taskrsrc = []
UNITS = {"SQM": "Square meter", "LM": "Linear meter"}


# ------------------------------------------------------------ calendar / dates
# working week + holidays are read from the calendar record copied from your P6 export
_days = {int(d): bool(re.search(r"\(0\|\|0\(s\|", body)) for d, body in
         re.findall(r"\(0\|\|(\d)\(\)\((.*?)\)\)\x7f\x7f(?=\s*\(0\|\||\))", CAL_ROW["clndr_data"] + "\x7f\x7f")}
WORKDAYS = {(n - 2) % 7 for n, w in _days.items() if w}      # P6 1=Sun..7=Sat -> python Mon=0..Sun=6
if not WORKDAYS: WORKDAYS = {0, 1, 2, 3, 4, 5}
HOLIDAYS = {(dt.date(1899, 12, 30) + dt.timedelta(days=int(n))) for n in re.findall(r"\(d\|(\d+)\)", CAL_ROW["clndr_data"])}

def is_wd(d):
    d = d.date() if isinstance(d, dt.datetime) else d
    return d.weekday() in WORKDAYS and d not in HOLIDAYS

def wd_count(a, b):
    n, d = 0, a.date()
    while d <= b.date():
        if is_wd(d): n += 1
        d += dt.timedelta(days=1)
    return max(n, 1)

def finish_after(start, ndays):
    d, n = start.date(), 0
    while True:
        if is_wd(d):
            n += 1
            if n >= ndays: return dt.datetime(d.year, d.month, d.day, DAY_END)
        d += dt.timedelta(days=1)

def next_wd(d):
    d = d.date() if isinstance(d, dt.datetime) else d
    while not is_wd(d): d += dt.timedelta(days=1)
    return dt.datetime(d.year, d.month, d.day, DAY_START)

def at(d, h): return dt.datetime(d.year, d.month, d.day, h) if d else None
def num(v):
    if isinstance(v, bool) or v is None: return None
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).strip().replace(",", ""))
    except ValueError: return None
def as_dt(v): return v if isinstance(v, dt.datetime) else None
def guid(key): return base64.b64encode(hashlib.md5(key.encode()).digest()).decode().rstrip("=")[:22]

def fmt(v):
    if v is None: return ""
    if isinstance(v, dt.datetime): return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else f"{v:.4f}".rstrip("0").rstrip(".")
    return str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ")

def table(name, rows):
    f = FIELDS[name]
    out = [f"%T\t{name}", "%F\t" + "\t".join(f)]
    for r in rows:
        extra = set(r) - set(f)
        assert not extra, (name, extra)
        out.append("%R\t" + "\t".join(r.get(k, "") if isinstance(r.get(k, ""), str) and k == "clndr_data" else fmt(r.get(k)) for k in f))
    return "\r\n".join(out)

# ------------------------------------------------------------------- read
wb = load_workbook(SRC, data_only=True); ws = wb["Sheet1"]
ws_fmt = load_workbook(SRC)["Sheet1"]
def is_red(cell):
    f = cell.fill
    return bool(f and f.fill_type and f.fgColor.type == "rgb" and f.fgColor.rgb == "FFFF0000")
buildings, excluded = OrderedDict(), OrderedDict()
for xl_row, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if r[0] in (None, ""): continue
    k = str(r[0]).strip()
    if is_red(ws_fmt.cell(row=xl_row, column=1)):
        excluded[k] = excluded.get(k, 0) + 1; continue
    buildings.setdefault(k, []).append((xl_row, r))
print("Excluded (red):", dict(excluded))

# ------------------------------------------------------------------ build
exceptions, activities_out, tasks, wbs_rows, udf_vals = [], [], [], [], []
def flag(code, name, xl_row, issue, action): exceptions.append((code, name, xl_row, issue, action))

W_ROOT, next_wbs, next_task = 14001, 14002, 56001
wbs_rows.append(dict(wbs_id=W_ROOT, proj_id=1, obs_id=OBS_ROW["obs_id"], seq_num=100, est_wt=1, proj_node_flag="Y",
                     sum_data_flag="N", status_code="WS_Open", wbs_short_name=PROJECT_ID, wbs_name=PROJECT_NAME,
                     parent_wbs_id=EPS_PARENT_ID, guid=guid("root")))
UDF = {"UOM": 101, "BQ": 102, "AQ": 103, "BMH": 104}
all_starts, all_ends = [], []

for bseq, (bname, rows) in enumerate(buildings.items()):
    bkey = re.sub(r"\s+", "_", bname)
    wb_id, wsup, wins = next_wbs, next_wbs + 1, next_wbs + 2; next_wbs += 3
    for wid, par, seq, short, nm in ((wb_id, W_ROOT, bseq, bname[:20], bname), (wsup, wb_id, 1, "SUPPLY", "SUPPLY / PRODUCTION"),
                                     (wins, wb_id, 2, "INSTALL", "INSTALLATION")):
        wbs_rows.append(dict(wbs_id=wid, proj_id=1, obs_id=OBS_ROW["obs_id"], seq_num=seq, est_wt=1, proj_node_flag="N",
                             sum_data_flag="N", status_code="WS_Open", wbs_short_name=short, wbs_name=nm,
                             parent_wbs_id=par, guid=guid(f"wbs{wid}{nm}")))
    seen = set()
    for xl_row, r in rows:
        desc = str(r[3]).strip(); typ = str(r[2]).strip().upper()
        sid = ID_MAP.get(desc)
        if sid is None:
            sid = ("SUP-9" if typ == "SUPPLY" else "INS-9") + str(len(seen) % 10)
            flag(f"{bkey}-{sid}", desc, xl_row, "Scope description not in the standard 30-item list", "Temporary ID assigned - add it to ID_MAP in the script")
        code = f"{bkey}-{sid}"
        assert len(code) <= 20, code
        if code in seen:
            flag(code, desc, xl_row, "Duplicate scope in the same building", "Row skipped"); continue
        seen.add(code)

        uom = r[4]; bls, blf, fs, ff = as_dt(r[5]), as_dt(r[6]), as_dt(r[7]), as_dt(r[8])
        bq, bmh, a_s, a_f, aq = num(r[9]), num(r[10]), as_dt(r[11]), as_dt(r[12]), num(r[14])
        tid = next_task; next_task += 1

        plan_s, plan_f = at(fs or bls, DAY_START), at(ff or blf, DAY_END)
        t_s, t_f = at(bls, DAY_START), at(blf, DAY_END)
        has_plan = plan_s is not None and plan_f is not None
        if has_plan and plan_f < plan_s:
            flag(code, desc, xl_row, "Planned finish earlier than planned start", "Planned finish set = planned start")
            plan_f = at(plan_s, DAY_END)
        orig_days = wd_count(plan_s, plan_f) if has_plan else NO_DATE_PLACEHOLDER_DAYS
        if not has_plan:
            flag(code, desc, xl_row, "No baseline/forecast dates in Excel", f"Placeholder {orig_days}-day duration, no planned dates - fill dates in Excel and re-run")

        pct = (aq / bq) if (bq and bq > 0 and aq is not None) else 0.0
        if pct > 1:
            flag(code, desc, xl_row, f"Actual qty ({aq:g}) exceeds budget qty ({bq:g}) = {pct:.0%}", "Physical % capped at 100")
            pct = 1.0

        bad_finish = bool(a_f and a_s and a_f < a_s)
        if a_f and not a_s:
            flag(code, desc, xl_row, "Actual finish without actual start", "Actual finish ignored"); a_f = None
        status = "TK_Complete" if (a_f and a_s and not bad_finish) else ("TK_Active" if a_s else "TK_NotStart")
        if bad_finish:
            flag(code, desc, xl_row, f"Actual finish {a_f:%d-%b-%Y} is before actual start {a_s:%d-%b-%Y}", "Actual finish NOT imported; left In Progress - please correct the date")
        if status == "TK_Complete" and pct < 0.995:
            flag(code, desc, xl_row, f"Has actual finish but quantity progress only {pct:.0%}", "Imported as Complete (100%)")
        if status == "TK_NotStart" and (aq or 0) > 0:
            flag(code, desc, xl_row, f"Actual qty {aq:g} recorded but no actual start date", "Imported as Not Started - add actual start in Excel")
        if status == "TK_Active" and not (bq and bq > 0):
            flag(code, desc, xl_row, "In progress but no budget quantity, so % cannot be calculated", "Physical % = 0")
        if status == "TK_Active" and pct >= 0.995 and not bad_finish:
            flag(code, desc, xl_row, "Quantity is 100% but no actual finish date", "Left In Progress, 1 day remaining - add actual finish")

        drtn = orig_days * HRS_PER_DAY
        t = {k: "" for k in FIELDS["TASK"]}
        t.update(task_id=tid, proj_id=1, wbs_id=(wsup if typ == "SUPPLY" else wins), clndr_id=CAL_ROW["clndr_id"],
                 rev_fdbk_flag="N", est_wt=1, lock_plan_flag="N", auto_compute_act_flag="N", complete_pct_type="CP_Phys",
                 task_type="TT_Task", duration_type="DT_FixedDUR2", status_code=status, task_code=code, task_name=desc,
                 act_work_qty=0, remain_work_qty=0, target_work_qty=0, target_equip_qty=0, act_equip_qty=0, remain_equip_qty=0,
                 target_drtn_hr_cnt=drtn, priority_type="PT_Normal", guid=guid(f"task{code}"), driving_path_flag="N",
                 act_this_per_work_qty=0, act_this_per_equip_qty=0, create_date=DATA_DATE, update_date=DATA_DATE,
                 create_user="NotPrmUser", update_user="NotPrmUser", target_start_date=t_s, target_end_date=t_f)
        if status == "TK_Complete":
            t.update(phys_complete_pct=100, remain_drtn_hr_cnt=0, act_start_date=at(a_s, DAY_START), act_end_date=at(a_f, DAY_END),
                     early_start_date=at(a_s, DAY_START), early_end_date=at(a_f, DAY_END),
                     late_start_date=at(a_s, DAY_START), late_end_date=at(a_f, DAY_END))
        elif status == "TK_Active":
            rem_days = max(1, math.ceil(orig_days * (1 - pct) - 1e-9))
            rs = next_wd(DATA_DATE); re_ = finish_after(rs, rem_days)
            t.update(phys_complete_pct=round(pct * 100, 2), remain_drtn_hr_cnt=rem_days * HRS_PER_DAY,
                     act_start_date=at(a_s, DAY_START), early_start_date=at(a_s, DAY_START), early_end_date=re_,
                     late_start_date=at(a_s, DAY_START), late_end_date=re_, restart_date=rs, reend_date=re_,
                     rem_late_start_date=rs, rem_late_end_date=re_)
        else:
            t.update(phys_complete_pct=0, remain_drtn_hr_cnt=drtn)
            if has_plan:
                t.update(early_start_date=plan_s, early_end_date=plan_f, late_start_date=plan_s, late_end_date=plan_f,
                         restart_date=plan_s, reend_date=plan_f, rem_late_start_date=plan_s, rem_late_end_date=plan_f,
                         cstr_type="CS_MSOA", cstr_date=plan_s)
                if plan_s < DATA_DATE:
                    flag(code, desc, xl_row, f"Not started but planned start {plan_s:%d-%b-%Y} is before data date", "Imported as is - P6 pushes it to the data date when scheduled (F9)")
            else:
                rs = next_wd(DATA_DATE); t.update(restart_date=rs, reend_date=finish_after(rs, orig_days))
        # ---- resource loading from the Excel quantities / man-hours
        item = re.sub(r"\s+(INSTALLATION|PRODUCTION|SUPPLY)$", "", desc)
        uom_s = str(uom).strip().upper() if uom else ""
        earn = 1.0 if status == "TK_Complete" else (pct if status == "TK_Active" else 0.0)
        asg = []                                           # (rsrc dict, target, actual, remain)
        lab = LABOR["INMHR" if typ == "INSTALLATION" else "PMHR"]
        lab_t = lab_a = lab_r = 0.0
        if bmh and bmh > 0:
            lab_t = round(bmh, 2); lab_a = round(bmh * earn, 2); lab_r = round(lab_t - lab_a, 2)
            asg.append((lab, lab_t, lab_a, lab_r))
        mat_t = bq if (bq and bq > 0) else 0.0
        if not mat_t and (aq or 0) > 0 and status != "TK_NotStart":
            flag(code, desc, xl_row, "Actual qty recorded but no budget qty", "Quantity resource target set equal to actual qty"); mat_t = aq
        mkey = None
        if mat_t > 0 and uom_s:
            mkey = f"{item}|{uom_s}"
            if mkey not in mat_rsrc:
                short = f"{ITEM_SHORT.get(item, re.sub(r'[^A-Z0-9]', '', item)[:10])}.{uom_s}"
                rid = 2001 + len(mat_rsrc)
                row = dict(MAT_RSRC_SAMPLE); row.update(rsrc_id=rid, rsrc_seq_num=200 + 10 * len(mat_rsrc), rsrc_name=f"{item} {uom_s}",
                           rsrc_short_name=short, guid=guid(f"rsrc{mkey}"), unit_id=(100 if uom_s == "SQM" else 101),
                           auto_compute_act_flag="N", clndr_id=CAL_ROW["clndr_id"])
                mat_rsrc[mkey] = row
            if status == "TK_NotStart": m_a = 0.0
            elif status == "TK_Complete": m_a = aq if aq is not None else mat_t
            else: m_a = aq or 0.0
            m_r = 0.0 if status == "TK_Complete" else max(0.0, mat_t - m_a)
            asg.append((mat_rsrc[mkey], mat_t, m_a, m_r))
        elif mat_t > 0:
            flag(code, desc, xl_row, "Budget qty present but no unit of measure", "No quantity resource assigned")
        rem_h = float(t["remain_drtn_hr_cnt"] or 0)
        for rr, tq, aq_, rq in asg:
            rate = LABOR_RATE.get(rr["rsrc_id"], 0.0) if rr["rsrc_type"] == "RT_Labor" else 0.0
            tr = {k: "" for k in FIELDS["TASKRSRC"]}
            tr.update(taskrsrc_id=15601 + len(taskrsrc), task_id=tid, proj_id=1, cost_qty_link_flag="Y", rsrc_id=rr["rsrc_id"],
                      remain_qty=rq, target_qty=tq, remain_qty_per_hr=(rq / rem_h if rem_h else 0), target_lag_drtn_hr_cnt=0,
                      target_qty_per_hr=tq / drtn, act_ot_qty=0, act_reg_qty=aq_, relag_drtn_hr_cnt=0, ot_factor="",
                      cost_per_qty=f"{rate:.4f}", target_cost=f"{tq * rate:.4f}", act_reg_cost=f"{aq_ * rate:.4f}", act_ot_cost="0.0000",
                      remain_cost=f"{rq * rate:.4f}", act_start_date=t["act_start_date"], act_end_date=t["act_end_date"],
                      restart_date=t["restart_date"], reend_date=t["reend_date"], target_start_date=t["target_start_date"],
                      target_end_date=t["target_end_date"], rollup_dates_flag="Y", ts_pend_act_end_flag="N",
                      guid=guid(f"tr{code}{rr['rsrc_short_name']}"), rate_type="COST_PER_QTY", act_this_per_cost="0.0000",
                      act_this_per_qty=0, rsrc_type=rr["rsrc_type"], cost_per_qty_source_type="ST_Rsrc",
                      create_user="NotPrmUser", create_date=DATA_DATE)
            taskrsrc.append(tr)
        t.update(act_work_qty=lab_a, remain_work_qty=lab_r, target_work_qty=lab_t, rsrc_id=(asg[0][0]["rsrc_id"] if asg else ""))
        tasks.append(t)
        for d in (plan_s, a_s): 
            if d: all_starts.append(d)
        for d in (plan_f, a_f, t.get("early_end_date") or None):
            if d: all_ends.append(d)
        udf_vals += [(UDF["UOM"], tid, None, str(uom) if uom else None), (UDF["BQ"], tid, bq, None),
                     (UDF["AQ"], tid, aq, None), (UDF["BMH"], tid, bmh, None)]
        activities_out.append((bname, code, desc, typ, uom, bq, aq, round(pct, 4), status, t_s, t_f, a_s,
                               a_f if status == "TK_Complete" else None, t["remain_drtn_hr_cnt"] / HRS_PER_DAY, xl_row,
                               lab_t or None, lab_a if lab_t else None, lab["rsrc_short_name"] if lab_t else "",
                               mat_rsrc[mkey]["rsrc_short_name"] if mkey else ""))

taskpred, link_report = [], []
code_to_tid = {t["task_code"]: t["task_id"] for t in tasks}
for bname in buildings:
    bkey = re.sub(r"\s+", "_", bname)
    for phase in ("SUP", "INS"):
        for p, q, typ_, lag in LINKS:
            pc, sc = f"{bkey}-{phase}-{p}", f"{bkey}-{phase}-{q}"
            if pc in code_to_tid and sc in code_to_tid:
                taskpred.append(dict(task_pred_id=90001 + len(taskpred), task_id=code_to_tid[sc], pred_task_id=code_to_tid[pc],
                                     proj_id=1, pred_proj_id=1, pred_type="PR_" + typ_, lag_hr_cnt=lag * HRS_PER_DAY,
                                     comments="", float_path="", aoa_flag="N"))
                link_report.append((bname, pc, sc, typ_, lag))
    for p, q, typ_, lag in CROSS_LINKS:
        pc, sc = f"{bkey}-{p}", f"{bkey}-{q}"
        if pc in code_to_tid and sc in code_to_tid:
            taskpred.append(dict(task_pred_id=90001 + len(taskpred), task_id=code_to_tid[sc], pred_task_id=code_to_tid[pc],
                                 proj_id=1, pred_proj_id=1, pred_type="PR_" + typ_, lag_hr_cnt=lag * HRS_PER_DAY,
                                 comments="", float_path="", aoa_flag="N"))
            link_report.append((bname, pc, sc, typ_, lag))

proj = dict(PROJECT_DEFAULTS)
proj.update(proj_id=1, proj_short_name=PROJECT_ID, clndr_id=CAL_ROW["clndr_id"], guid=guid("project"),
            last_recalc_date=DATA_DATE, plan_start_date=min(all_starts + [DATA_DATE]), scd_end_date=max(all_ends),
            add_date=DATA_DATE, task_code_prefix="", task_code_prefix_flag="N", wbs_max_sum_level=2, add_by_name="admin",
            last_tasksum_date="", fcst_start_date="")
for k in ("last_checksum", "last_fin_dates_id", "last_baseline_update_date", "new_fin_dates_id", "next_data_date",
          "close_period_flag", "sum_refresh_date", "trsrcsum_loaded", "baselines_to_export", "baseline_names_to_export"):
    proj[k] = ""
for k in ("plan_start_date", "scd_end_date", "last_recalc_date", "add_date"):
    proj[k] = fmt(proj[k])

# --------------------------------------------------------------- assemble
def build_xer(with_res, with_logic=True):
    parts = ["\t".join(["ERMHDR", P6_VERSION, DATA_DATE.strftime("%Y-%m-%d"), "Project", "admin", "admin",
                        "dbxDatabaseNoName", "Project Management", "USD"])]
    parts.append(table("CURRTYPE", [USD_ROW]))
    parts.append(table("OBS", [OBS_ROW]))
    if with_res:
        used = {"SQM", "LM"} & {k.split("|")[1] for k in mat_rsrc}
        parts.append(table("UMEASURE", [dict(unit_id=100 if u == "SQM" else 101, seq_num=0 if u == "SQM" else 1,
                                            unit_abbrev=u, unit_name=UNITS[u]) for u in sorted(used, reverse=True)]))
    parts.append(table("PROJECT", [proj]))
    parts.append(table("CALENDAR", [CAL_ROW]))
    parts.append(table("PROJWBS", wbs_rows))
    if with_res:
        parts.append(table("RSRC", list(LABOR_RSRC_ROWS) + list(mat_rsrc.values())))
        rates = list(LABOR_RATE_ROWS)
        for n, r in enumerate(mat_rsrc.values()):
            rates.append(dict(rsrc_rate_id=1900 + n, rsrc_id=r["rsrc_id"], max_qty_per_hr=1, cost_per_qty="0.0000",
                              start_date="2025-01-01 00:00", cost_per_qty2="0.0000", cost_per_qty3="0.0000",
                              cost_per_qty4="0.0000", cost_per_qty5="0.0000"))
        parts.append(table("RSRCRATE", rates))
    if with_res:
        parts.append(table("TASK", tasks))
        if with_logic: parts.append(table("TASKPRED", taskpred))
        parts.append(table("TASKRSRC", taskrsrc))
    else:
        plain = [dict(t, rsrc_id="", act_work_qty=0, remain_work_qty=0, target_work_qty=0) for t in tasks]
        parts.append(table("TASK", plain))
    return "\r\n".join(parts) + "\r\n%E\r\n"

os.makedirs(OUT, exist_ok=True)
for name, logic in (("AZAL_PROGRAM_P6v17_resource_loaded.xer", True), ("AZAL_PROGRAM_P6v17_resource_loaded_NO_LOGIC.xer", False)):
    with open(os.path.join(OUT, name), "w", encoding="cp1252", newline="") as f:
        f.write(build_xer(True, logic))

# ------------------------------------------------------------------ report
rep = Workbook()
hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
def sheet(w, headers, data, widths):
    w.append(headers)
    for c in w[1]: c.font, c.fill = hf, hfill
    for d in data: w.append(list(d))
    for i, x in enumerate(widths, 1): w.column_dimensions[get_column_letter(i)].width = x
    w.freeze_panes = "A2"; w.auto_filter.ref = w.dimensions
s1 = rep.active; s1.title = "Summary"
by = {}
for a in activities_out: by.setdefault(a[0], Counter())[a[8]] += 1
exc_by = Counter(e[0].rsplit("-", 2)[0] for e in exceptions)
sheet(s1, ["Building (WBS)", "Activities", "Not Started", "In Progress", "Complete", "Exceptions"],
      [(p, sum(c.values()), c["TK_NotStart"], c["TK_Active"], c["TK_Complete"], exc_by.get(re.sub(r"\s+", "_", p), 0)) for p, c in by.items()],
      [22, 12, 13, 13, 11, 12])
sheet(rep.create_sheet("Exceptions"), ["Activity ID", "Activity Name", "Excel row", "Issue", "How it was imported"], exceptions, [24, 40, 10, 70, 70])
s3 = rep.create_sheet("Activity Map")
sheet(s3, ["Building", "Activity ID", "Activity Name", "Type", "UOM", "Budget Qty", "Actual Qty", "Phys % (0-1)", "P6 Status",
           "Planned Start", "Planned Finish", "Actual Start", "Actual Finish", "Remaining Days", "Excel row", "Budget Mhr", "Earned Mhr", "Labor Resource", "Qty Resource"],
      activities_out, [18, 20, 40, 14, 7, 11, 11, 11, 13, 13, 13, 13, 13, 12, 10, 12, 12, 14, 18])
for r in s3.iter_rows(min_row=2):
    for c in r[9:13]: c.number_format = "dd-mmm-yyyy"
    r[7].number_format = "0.0%"
sheet(rep.create_sheet("Relationships"), ["Building", "Predecessor", "Successor", "Type", "Lag (working days)"], link_report, [20, 24, 24, 8, 18])
rep.save(os.path.join(OUT, "AZAL_XER_import_report.xlsx"))
print(f"relationships={len(taskpred)}")
print(f"resources={len(LABOR_RSRC_ROWS)+len(mat_rsrc)} assignments={len(taskrsrc)}")
print(f"buildings={len(buildings)} wbs={len(wbs_rows)} tasks={len(tasks)} exceptions={len(exceptions)}")
print(Counter(t["status_code"] for t in tasks), "workdays", sorted(WORKDAYS), "holidays", len(HOLIDAYS))
