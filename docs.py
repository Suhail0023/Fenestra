# FenOps - Copyright (c) 2026 Charles Couteau. All rights reserved.
# Proprietary and confidential. Unauthorised copying, modification, distribution or reverse engineering is prohibited. See LICENSE.txt.
"""FenOps document engine.

  * serial numbering of every official letter / form, with a custom format and start number per type
  * the document register (a stored copy of each item, with status, sent flag and file)
  * customisable formats: Word (.docx), Excel (.xlsx) or PDF templates are filled with the item's data
  * reports and downloads (Excel, PDF, CSV, ZIP of copies)

Everything is stored in the app's SQLite file. Libraries:
  Word templates and the default Word format : Python standard library only
  Excel templates and Excel reports          : openpyxl   (pip install openpyxl)
  PDF templates, PDF reports                 : pypdf, reportlab   (pip install pypdf reportlab)
  "Download as PDF" of Word/Excel copies     : LibreOffice, if installed (otherwise a built-in PDF layout is used)
"""
import csv, datetime as dt, io, json, os, re, shutil, sqlite3, struct, subprocess, tempfile, zipfile
from xml.sax.saxutils import escape as xesc, unescape as xunesc

DB_FILE = None
COMPANY = lambda: {}
PROJECT_NAME = lambda code: code
LOGO_PATH = lambda: None
MAX_TEMPLATE = 8 * 1024 * 1024
COMMERCIAL = {"IPA", "VOP", "VO", "DLY", "EOT", "INS"}   # only Management and the project manager see these in the register

# --------------------------------------------------------------------------------------------
# Document types and the fields a template can use
# --------------------------------------------------------------------------------------------
STD_FIELDS = [("number", "Serial number, e.g. WIR-001"), ("date", "Issue date"), ("project", "Project name"), ("project_code", "Project code"),
              ("company", "Your company name"), ("company_address", "Company address"), ("company_phone", "Company phone"), ("company_email", "Company email"),
              ("to", "Addressee"), ("from", "Sender"), ("subject", "Subject / title"), ("ref", "Related record (task, instruction, delivery ...)"),
              ("status", "Status"), ("revision", "Revision number"), ("issued_by", "Issued by (user)"), ("body", "Full text of the letter or form"),
              ("generated_on", "Date this copy was generated")]
DOC_TYPES = {
 "WIR": dict(label="Work inspection request (WIR)", pattern="WIR-{NNN}",
             fields=[("task", "Activity / task"), ("quantity", "Quantity inspected"), ("uom", "Unit"), ("requested_on", "Requested on"), ("decision", "Approved / Rejected / Pending"),
                     ("decided_on", "Decision date"), ("decided_by", "Decided by"), ("notes", "Inspector notes"), ("scope", "Inspection scope"), ("package", "WIR package number")],
             lists={"checks": [("item", "Checklist item"), ("result", "Pass / Fail / value")]}),
 "WIRP": dict(label="WIR package / transmittal", pattern="WIRP-{NNN}",
              fields=[("package_note", "Cover note"), ("wir_count", "Number of WIRs"), ("total_quantity", "Total quantity"), ("decision", "Decision"), ("decided_on", "Decision date")],
              lists={"wirs": [("number", "WIR number"), ("task", "Task"), ("quantity", "Quantity")]}),
 "MIR": dict(label="Material inspection request (MIR)", pattern="MIR-{NNN}",
             fields=[("material", "Material"), ("supplier", "Supplier"), ("delivery_note", "Delivery note"), ("quantity", "Quantity"), ("uom", "Unit"), ("location", "Storage location"),
                     ("delivery_date", "Delivery date"), ("decision", "Decision"), ("decided_on", "Decision date"), ("notes", "Notes")],
             lists={"checks": [("item", "Checklist item"), ("ok", "Yes / No")]}),
 "IPA": dict(label="Interim payment application (IPA)", pattern="IPA-{NN}",
             fields=[("month", "Month"), ("cutoff", "Cut-off date"), ("gross", "Gross value to date"), ("previous_gross", "Previously certified gross"), ("this_period", "Gross this period"),
                     ("retention_pct", "Retention %"), ("retention", "Retention amount"), ("net", "Net amount due"), ("currency", "Currency"),
                     ("mir_pct", "Claim % materials (MIR)"), ("wir_pct", "Claim % installed (WIR)"), ("snag_pct", "Claim % snag IR")],
             lists={"items": [("name", "BOQ line"), ("unit", "Unit"), ("qty", "Contract qty"), ("rate", "Rate"), ("mir_qty", "MIR-approved qty"), ("wir_qty", "WIR-approved qty"),
                              ("snag_qty", "Snag-cleared qty"), ("pct", "% of line"), ("previous", "Previous amount"), ("cumulative", "Cumulative amount"), ("current", "This period")]}),
 "DLY": dict(label="Delay report / claim letter", pattern="DLY-{NNN}",
             fields=[("month", "Month"), ("eot_days", "Extension of time claimed (days)"), ("delay_events", "Number of delay events"), ("cost", "Cost claimed"), ("fidic_edition", "FIDIC edition")],
             lists={}),
 "EOT": dict(label="Extension of time (EOT) letter", pattern="EOT-{NNN}",
             fields=[("days", "Days late"), ("baseline_finish", "Baseline completion"), ("forecast_finish", "Forecast completion"), ("cause", "Cause")], lists={}),
 "INS": dict(label="Instruction record and acknowledgement", pattern="INS-{NNN}",
             fields=[("instruction_date", "Instruction date"), ("instruction_from", "Given by"), ("instruction_ref", "Main contractor reference"), ("instruction_kind", "Site / Engineering"),
                     ("description", "Description"), ("quantity", "Estimated quantity"), ("deadline", "Details due by")], lists={}),
 "VOP": dict(label="Variation proposal letter", pattern="VOP-{NNN}",
             fields=[("instruction_no", "Instruction number"), ("quantity", "Quantity"), ("rate", "Rate"), ("value", "Proposed value"), ("time_impact", "Time impact (days)")], lists={}),
 "VO": dict(label="Variation order register entry", pattern="VO-{NNN}",
            fields=[("description", "Description"), ("quantity", "Quantity"), ("value", "Value")], lists={}),
 "TRQ": dict(label="Store tools requisition", pattern="TRQ-{NNN}",
             fields=[("date_required", "Date required"), ("requested_by", "Requested by")],
             lists={"tools": [("employee_id", "Employee ID"), ("name", "Employee"), ("tool", "Tool"), ("qty", "Quantity")]}),
 "WARN": dict(label="Warning letter", pattern="W{NNNN}",
              fields=[("employee_id", "Employee ID"), ("employee", "Employee"), ("job_role", "Role"), ("team", "Team"), ("task", "Task"), ("warning_level", "First / Second / Final")], lists={}),
 "DN": dict(label="Danger notice", pattern="DN-{NNN}",
            fields=[("location", "Location"), ("severity", "Severity"), ("hazard", "Hazard"), ("action_required", "Immediate action"), ("raised_by", "Raised by")], lists={}),
 "PTW": dict(label="Permit to work", pattern="PTW-{NNN}",
             fields=[("permit_type", "Permit type"), ("valid_from", "Valid from"), ("valid_to", "Valid to"), ("task", "Task"), ("requested_by", "Requested by")], lists={}),
 "SIR": dict(label="Snag / de-snag inspection approval", pattern="SIR-{NNN}", fields=[("task", "Task")], lists={}),
 "SN": dict(label="Snag item", pattern="SN-{NNN}",
            fields=[("location", "Location"), ("description", "Description"), ("owner", "Responsible"), ("due", "Due date")], lists={}),
 "DEL": dict(label="Material delivery record", pattern="DEL-{NNN}",
             fields=[("supplier", "Supplier"), ("delivery_note", "Delivery note"), ("material", "Material"), ("quantity", "Quantity"), ("uom", "Unit")], lists={}),
 "SOR": dict(label="Site observation report (rejected material)", pattern="SOR-{NNN}",
             fields=[("material", "Material"), ("quantity", "Quantity"), ("uom", "Unit"), ("supplier", "Supplier"), ("delivery_note", "Delivery note"), ("mir_no", "MIR number"),
                     ("rejection_reason", "Reason for rejection"), ("action_required", "Action required"), ("location", "Stored at"), ("raised_by", "Raised by"), ("status", "Status")], lists={}),
 "CERT": dict(label="Certificate of recognition", pattern="CERT-{NNN}",
              fields=[("recipient", "Recipient"), ("category", "Award"), ("month", "Month"), ("score", "KPI score"), ("rank", "Rank"), ("approved_by", "Approved by"), ("approved_on", "Approved on"), ("citation", "Citation")], lists={}),
 "LTR": dict(label="General letter / other official form", pattern="LTR-{NNN}", fields=[], lists={}),
}
MONEY_KEYS = {"gross", "previous_gross", "this_period", "retention", "net", "value", "cost", "rate", "previous", "cumulative", "current"}

def type_info(t):
    d = DOC_TYPES[t]
    return {"type": t, "label": d["label"], "default_pattern": d["pattern"], "field_help": [[n, h] for n, h in STD_FIELDS + d["fields"]],
            "lists": {k: [[c, h] for c, h in v] for k, v in d["lists"].items()}}

# --------------------------------------------------------------------------------------------
# Database
# --------------------------------------------------------------------------------------------
def init(db_file, company_fn=None, project_name_fn=None, logo_fn=None):
    global DB_FILE, COMPANY, PROJECT_NAME, LOGO_PATH
    DB_FILE = db_file
    if company_fn: COMPANY = company_fn
    if project_name_fn: PROJECT_NAME = project_name_fn
    if logo_fn: LOGO_PATH = logo_fn
    c = conn()
    try:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS documents (id INTEGER PRIMARY KEY, project TEXT NOT NULL, type TEXT NOT NULL, number TEXT NOT NULL, rev INTEGER NOT NULL DEFAULT 0, seq INTEGER,
            date TEXT, title TEXT, to_party TEXT, from_party TEXT, ref TEXT, status TEXT NOT NULL DEFAULT 'Reserved', void_reason TEXT, sent INTEGER DEFAULT 0,
            body TEXT, data TEXT, src TEXT, sig TEXT, dup INTEGER DEFAULT 0, created_at TEXT, created_by TEXT, updated_at TEXT, updated_by TEXT,
            UNIQUE (project, type, number, rev));
        CREATE INDEX IF NOT EXISTS ix_documents_proj ON documents(project, type, date);
        CREATE TABLE IF NOT EXISTS doc_files (doc_id INTEGER PRIMARY KEY, filename TEXT, mime TEXT, blob BLOB, tpl_name TEXT, made_at TEXT, error TEXT);
        CREATE TABLE IF NOT EXISTS doc_counters (project TEXT NOT NULL, type TEXT NOT NULL, period TEXT NOT NULL, next INTEGER NOT NULL, PRIMARY KEY (project, type, period));
        CREATE TABLE IF NOT EXISTS doc_numbering (project TEXT NOT NULL, type TEXT NOT NULL, pattern TEXT, start INTEGER, reset TEXT, PRIMARY KEY (project, type));
        CREATE TABLE IF NOT EXISTS doc_templates (id INTEGER PRIMARY KEY, project TEXT NOT NULL DEFAULT '', type TEXT NOT NULL, name TEXT, filename TEXT, kind TEXT, blob BLOB,
            opts TEXT, active INTEGER DEFAULT 1, uploaded_by TEXT, uploaded_at TEXT, size INTEGER);
        CREATE INDEX IF NOT EXISTS ix_doc_templates ON doc_templates(type, project);
        """)
    finally: c.close()

def conn():
    c = sqlite3.connect(DB_FILE, timeout=30, isolation_level=None)
    c.execute("PRAGMA journal_mode=WAL"); c.execute("PRAGMA synchronous=FULL"); c.row_factory = sqlite3.Row
    return c

def now(): return dt.datetime.now().strftime("%Y-%m-%d %H:%M")
def today(): return dt.date.today().strftime("%Y-%m-%d")

# --------------------------------------------------------------------------------------------
# Numbering
# --------------------------------------------------------------------------------------------
def fmt_number(pattern, n, project, date, type_=""):
    date = (date or today())[:10]; y, m, d = (date.split("-") + ["01", "01"])[:3]
    def sub(mm):
        tok = mm.group(1)
        if set(tok) == {"N"}: return str(n).zfill(len(tok))
        return {"YYYY": y, "YY": y[2:], "MM": m, "DD": d, "PRJ": project, "TYPE": type_}.get(tok, mm.group(0))
    return re.sub(r"\{([A-Za-z]+)\}", sub, pattern)

def check_pattern(p):
    p = (p or "").strip()
    if not re.search(r"\{N+\}", p): raise ValueError("The format must contain a number token such as {NNN}, e.g. WIR-{NNN} or {PRJ}/WIR/{YYYY}/{NNN}")
    if len(p) > 60 or re.search(r"[\x00-\x1f\\\"<>]", p): raise ValueError("The format is too long or has characters that are not allowed")
    bad = [t for t in re.findall(r"\{([A-Za-z]+)\}", p) if not (set(t) == {"N"} or t in ("YYYY", "YY", "MM", "DD", "PRJ", "TYPE"))]
    if bad: raise ValueError("Unknown token {%s}. Use {NNN}, {YYYY}, {YY}, {MM}, {DD}, {PRJ}, {TYPE}" % bad[0])
    return p

def period_key(reset, date):
    date = (date or today())[:10]
    return date[:4] if reset == "yearly" else date[:7] if reset == "monthly" else ""

def get_cfg(c, project, type_):
    r = c.execute("SELECT pattern, start, reset FROM doc_numbering WHERE project=? AND type=?", (project, type_)).fetchone()
    d = DOC_TYPES[type_]
    return {"pattern": (r["pattern"] if r and r["pattern"] else d["pattern"]), "start": int(r["start"]) if r and r["start"] is not None else 1,
            "reset": (r["reset"] if r and r["reset"] else "none")}

def _free_number(c, project, type_, cfg, date, n):
    while True:
        num = fmt_number(cfg["pattern"], n, project, date, type_)
        if not c.execute("SELECT 1 FROM documents WHERE project=? AND type=? AND number=?", (project, type_, num)).fetchone(): return n, num
        n += 1

def _base(c, project, type_, cfg, per, floor=0):
    r = c.execute("SELECT next FROM doc_counters WHERE project=? AND type=? AND period=?", (project, type_, per)).fetchone()
    return max(int(r["next"]), 1) if r else max(cfg["start"], int(floor or 0), 1)

def peek(c, project, type_, date=None, floor=0):
    cfg = get_cfg(c, project, type_)
    return _free_number(c, project, type_, cfg, date, _base(c, project, type_, cfg, period_key(cfg["reset"], date), floor))

def allocate(project, type_, date=None, by="", floor=0):
    """Take the next serial number. A 'Reserved' row is written at once so a number is never lost or issued twice."""
    if type_ not in DOC_TYPES: raise ValueError("unknown document type")
    date = (date or today())[:10]; c = conn()
    try:
        c.execute("BEGIN IMMEDIATE")
        try:
            cfg = get_cfg(c, project, type_); per = period_key(cfg["reset"], date)
            n, num = _free_number(c, project, type_, cfg, date, _base(c, project, type_, cfg, per, floor))
            c.execute("INSERT OR REPLACE INTO doc_counters (project, type, period, next) VALUES (?,?,?,?)", (project, type_, per, n + 1))
            cur = c.execute("INSERT INTO documents (project, type, number, rev, seq, date, status, created_at, created_by, updated_at, updated_by) VALUES (?,?,?,?,?,?,'Reserved',?,?,?,?)",
                            (project, type_, num, 0, n, date, now(), by, now(), by))
            c.execute("COMMIT")
        except Exception: c.execute("ROLLBACK"); raise
        return {"id": cur.lastrowid, "number": num, "seq": n, "type": type_, "next": n + 1, "pattern": cfg["pattern"], "reset": cfg["reset"]}
    finally: c.close()

def numbering_config(project):
    c = conn()
    try:
        out = []
        for t in DOC_TYPES:
            cfg = get_cfg(c, project, t); n, num = peek(c, project, t)
            used = c.execute("SELECT COUNT(*) FROM documents WHERE project=? AND type=? AND status<>'Reserved'", (project, t)).fetchone()[0]
            res = c.execute("SELECT COUNT(*) FROM documents WHERE project=? AND type=? AND status='Reserved'", (project, t)).fetchone()[0]
            out.append(dict(type_info(t), pattern=cfg["pattern"], start=cfg["start"], reset=cfg["reset"], next=n, next_number=num, issued=used, reserved=res))
        return out
    finally: c.close()

def set_numbering(project, type_, pattern, start, reset, next_=None):
    if type_ not in DOC_TYPES: raise ValueError("unknown document type")
    pattern = check_pattern(pattern); reset = reset if reset in ("none", "yearly", "monthly") else "none"
    start = max(1, int(start or 1)); c = conn()
    try:
        c.execute("BEGIN IMMEDIATE")
        try:
            old = get_cfg(c, project, type_); had = c.execute("SELECT 1 FROM doc_numbering WHERE project=? AND type=?", (project, type_)).fetchone()
            c.execute("INSERT OR REPLACE INTO doc_numbering (project, type, pattern, start, reset) VALUES (?,?,?,?,?)", (project, type_, pattern, start, reset))
            per = period_key(reset, None)
            # a changed start number takes effect for the next item issued
            if next_ not in (None, ""): tgt = int(next_)
            elif (not had and start != 1) or (had and start != old["start"]): tgt = start
            else: tgt = None
            if tgt is not None: c.execute("INSERT OR REPLACE INTO doc_counters (project, type, period, next) VALUES (?,?,?,?)", (project, type_, per, max(1, tgt)))
            c.execute("COMMIT")
        except Exception: c.execute("ROLLBACK"); raise
    finally: c.close()

# --------------------------------------------------------------------------------------------
# Register
# --------------------------------------------------------------------------------------------
def register(project, type_, d, by=""):
    """Create or update the stored copy of an official document and (re)build its file.
    d: number, rev, date, title, to, from, ref, status, sent, body, src, sig, fields{}, lists{}"""
    if type_ not in DOC_TYPES: raise ValueError("unknown document type")
    number = str(d.get("number") or "").strip()
    if not number: raise ValueError("number required")
    rev = int(d.get("rev") or 0); sig = d.get("sig") or ""; src = d.get("src") or ""
    data = json.dumps({"fields": d.get("fields") or {}, "lists": d.get("lists") or {}}, ensure_ascii=False)
    c = conn(); changed = True
    try:
        c.execute("BEGIN IMMEDIATE")
        try:
            row = c.execute("SELECT * FROM documents WHERE project=? AND type=? AND number=? AND rev=?", (project, type_, number, rev)).fetchone()
            dup = 0
            if row and src and row["src"] and row["src"] != src and row["status"] != "Reserved":
                # the same number was used for a different record (typically created offline on two devices): keep both, mark the later one
                base = number; k = 0
                while True:
                    k += 1; number = base + chr(64 + k) if k <= 26 else base + "-" + str(k)
                    row = c.execute("SELECT * FROM documents WHERE project=? AND type=? AND number=? AND rev=?", (project, type_, number, rev)).fetchone()
                    if not row or row["src"] == src: break
                dup = 1
            status = d.get("status") or "Issued"
            if row and row["status"] == "Void": status = "Void"
            vals = (d.get("date") or (row["date"] if row else today()), str(d.get("title") or "")[:300], str(d.get("to") or "")[:200], str(d.get("from") or "")[:200],
                    str(d.get("ref") or "")[:200], status, 1 if d.get("sent") else 0, d.get("body") or "", data, src, sig)
            if row:
                changed = (row["sig"] or "") != sig or not sig
                c.execute("UPDATE documents SET date=?, title=?, to_party=?, from_party=?, ref=?, status=?, sent=?, body=?, data=?, src=?, sig=?, updated_at=?, updated_by=? WHERE id=?",
                          vals + (now(), by, row["id"])); did = row["id"]
                if dup: c.execute("UPDATE documents SET dup=1 WHERE id=?", (did,))
            else:
                m = re.search(r"(\d+)\D*$", number); seq = int(m.group(1)) if m else None
                cur = c.execute("INSERT INTO documents (date,title,to_party,from_party,ref,status,sent,body,data,src,sig,project,type,number,rev,seq,dup,created_at,created_by,updated_at,updated_by) "
                                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", vals + (project, type_, number, rev, seq, dup, now(), by, now(), by)); did = cur.lastrowid
            c.execute("COMMIT")
        except Exception: c.execute("ROLLBACK"); raise
        has = c.execute("SELECT 1 FROM doc_files WHERE doc_id=? AND LENGTH(blob)>0", (did,)).fetchone()
    finally: c.close()
    if changed or not has: make_file(project, did)
    return {"id": did, "number": number, "dup": bool(dup)}

def _row_ctx(row):
    data = json.loads(row["data"] or "{}"); f = dict(data.get("fields") or {}); co = COMPANY() or {}
    base = {"number": row["number"], "date": row["date"] or "", "project": PROJECT_NAME(row["project"]), "project_code": row["project"], "company": co.get("name", ""),
            "company_address": co.get("address", ""), "company_phone": co.get("phone", ""), "company_email": co.get("email", ""), "to": row["to_party"] or "",
            "from": row["from_party"] or "", "subject": row["title"] or "", "ref": row["ref"] or "", "status": row["status"], "revision": row["rev"],
            "issued_by": row["created_by"] or "", "body": row["body"] or "", "generated_on": today()}
    for k, v in base.items():
        if k in ("number", "status", "revision") or f.get(k) in (None, ""): f[k] = v
    return {"fields": f, "lists": data.get("lists") or {}, "type": row["type"], "_unknown": set()}

def active_template(c, project, type_):
    return c.execute("SELECT * FROM doc_templates WHERE type=? AND active=1 AND project IN (?, '') ORDER BY (project=?) DESC, id DESC LIMIT 1", (type_, project, project)).fetchone()

def make_file(project, doc_id):
    """Render and store the copy using the active uploaded format for the type (or the built-in default)."""
    c = conn()
    try:
        row = c.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not row: return
        tpl = active_template(c, project, row["type"]); ctx = _row_ctx(row); err = ""; used = ""
        try:
            fn, mime, blob = render(row["type"], ctx, tpl); used = tpl["name"] if tpl else ""
        except Exception as e:
            err = "%s: %s" % (type(e).__name__, e)
            try: fn, mime, blob = render(row["type"], _row_ctx(row), None)
            except Exception as e2: fn, mime, blob, err = "", "", b"", err + " / " + str(e2)
        fn = safe_name(row["number"]) + os.path.splitext(fn)[1] if fn else ""
        c.execute("INSERT OR REPLACE INTO doc_files (doc_id, filename, mime, blob, tpl_name, made_at, error) VALUES (?,?,?,?,?,?,?)", (doc_id, fn, mime, blob, used, now(), err))
    finally: c.close()

def safe_name(s): return re.sub(r"[^A-Za-z0-9._-]+", "_", str(s)).strip("_") or "document"

def list_docs(project, type_="", q="", status="", d_from="", d_to="", limit=5000, hide=()):
    c = conn()
    try:
        sql = ("SELECT d.id, d.type, d.number, d.rev, d.date, d.title, d.to_party, d.from_party, d.ref, d.status, d.void_reason, d.sent, d.created_by, d.created_at, d.updated_at, d.dup, d.data, "
               "(SELECT COUNT(*) FROM doc_files f WHERE f.doc_id=d.id AND LENGTH(f.blob)>0) AS has_file, (SELECT tpl_name FROM doc_files f WHERE f.doc_id=d.id) AS tpl_name, "
               "(SELECT error FROM doc_files f WHERE f.doc_id=d.id) AS file_error FROM documents d WHERE d.project=?")
        a = [project]
        if type_: sql += " AND d.type=?"; a.append(type_)
        if status: sql += " AND d.status=?"; a.append(status)
        if d_from: sql += " AND d.date>=?"; a.append(d_from)
        if d_to: sql += " AND d.date<=?"; a.append(d_to)
        if q: sql += " AND (d.number LIKE ? OR d.title LIKE ? OR d.to_party LIKE ? OR d.ref LIKE ? OR d.body LIKE ?)"; a += ["%" + q + "%"] * 5
        sql += " ORDER BY d.date DESC, d.id DESC LIMIT ?"; a.append(limit)
        out = []
        for r in c.execute(sql, a):
            x = dict(r)
            if x["type"] in hide: continue
            try: f = (json.loads(x.pop("data") or "{}").get("fields") or {})
            except Exception: f = {}
            x["amount"] = _key_figure(x["type"], f); x["label"] = DOC_TYPES.get(x["type"], {}).get("label", x["type"]); out.append(x)
        return out
    finally: c.close()

def _key_figure(t, f):
    try:
        if t == "IPA": return "%s %s" % (f.get("currency", ""), format(round(float(f.get("net") or 0)), ","))
        if t == "VOP": return "%s" % format(round(float(f.get("value") or 0)), ",")
        if t == "DLY": return "%s day(s)" % f.get("eot_days", 0)
        if t == "EOT": return "%s day(s)" % f.get("days", "")
        if t in ("WIR", "MIR", "DEL", "WIRP"): return "%s %s" % (f.get("quantity", f.get("total_quantity", "")), f.get("uom", ""))
    except Exception: pass
    return ""

def get_doc(project, doc_id):
    c = conn()
    try:
        r = c.execute("SELECT * FROM documents WHERE id=? AND project=?", (doc_id, project)).fetchone()
        if not r: return None
        x = dict(r); x["data"] = json.loads(x.get("data") or "{}"); x["label"] = DOC_TYPES.get(x["type"], {}).get("label", x["type"])
        f = c.execute("SELECT filename, mime, tpl_name, made_at, error, LENGTH(blob) AS size FROM doc_files WHERE doc_id=?", (doc_id,)).fetchone()
        x["file"] = dict(f) if f else None
        return x
    finally: c.close()

def doc_type_of(project, doc_id):
    c = conn()
    try:
        r = c.execute("SELECT type FROM documents WHERE id=? AND project=?", (doc_id, project)).fetchone()
        return r["type"] if r else None
    finally: c.close()

def sigs(project):
    c = conn()
    try: return {"%s|%s|%s" % (r["type"], r["number"], r["rev"]): r["sig"] or "" for r in c.execute("SELECT type, number, rev, sig FROM documents WHERE project=? AND status<>'Reserved'", (project,))}
    finally: c.close()

def get_file(project, doc_id, regen=False, as_=""):
    """(filename, mime, bytes) of the stored copy. regen re-renders with the format that is active now. as_='pdf' converts."""
    c = conn()
    try:
        r = c.execute("SELECT * FROM documents WHERE id=? AND project=?", (doc_id, project)).fetchone()
        if not r: return None
        f = c.execute("SELECT filename, mime, blob FROM doc_files WHERE doc_id=?", (doc_id,)).fetchone()
    finally: c.close()
    if regen or not f or not f["blob"]:
        make_file(project, doc_id)
        c = conn()
        try: f = c.execute("SELECT filename, mime, blob FROM doc_files WHERE doc_id=?", (doc_id,)).fetchone()
        finally: c.close()
    if not f or not f["blob"]: return None
    fn, mime, blob = f["filename"], f["mime"], bytes(f["blob"])
    if as_ == "pdf" and not fn.lower().endswith(".pdf"):
        pdf = to_pdf(blob, os.path.splitext(fn)[1])
        if pdf is None:
            c = conn()
            try: ctx = _row_ctx(r)
            finally: c.close()
            pdf = default_pdf(r["type"], ctx)
        return os.path.splitext(fn)[0] + ".pdf", "application/pdf", pdf
    return fn, mime, blob

def void_doc(project, doc_id, reason, by):
    c = conn()
    try: c.execute("UPDATE documents SET status='Void', void_reason=?, updated_at=?, updated_by=? WHERE id=? AND project=?", (str(reason or "")[:300], now(), by, doc_id, project))
    finally: c.close()

# --------------------------------------------------------------------------------------------
# Value helpers and token lookup
# --------------------------------------------------------------------------------------------
TOK = re.compile(r"\{\{\s*([A-Za-z0-9_.\- ]+?)\s*\}\}")
def _norm(k): return re.sub(r"[\s\-]+", "_", str(k).strip().lower())
_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
def clean_xml(s): return _BAD_XML.sub("", str(s))

def fmt_val(v, key=""):
    if v is None: return ""
    if isinstance(v, bool): return "Yes" if v else "No"
    if isinstance(v, (int, float)):
        k = key.split(".")[-1]
        if k in MONEY_KEYS: return format(round(v), ",") if abs(v) >= 100 or float(v).is_integer() else format(round(v, 2), ",.2f")
        if isinstance(v, float): return ("%.2f" % v).rstrip("0").rstrip(".")
        return str(v)
    return str(v)

def lookup(ctx, key, item=None):
    k = _norm(key)
    if "." in k:
        ln, col = k.split(".", 1)
        if item is not None: return item.get(col, item.get(col.replace("_", ""), ""))
        return ""
    f = {_norm(a): b for a, b in ctx["fields"].items()}
    if k in f: return f[k]
    ctx["_unknown"].add(key.strip()); return ""

def tokens_in(text): return [m.group(1) for m in TOK.finditer(text or "")]

def known_tokens(type_):
    d = DOC_TYPES[type_]; s = {n for n, _ in STD_FIELDS + d["fields"]}
    for ln, cols in d["lists"].items(): s |= {"%s.%s" % (ln, c) for c, _ in cols}
    return s

def sample_ctx(type_):
    d = DOC_TYPES[type_]; co = COMPANY() or {}; f = {}
    for n, h in STD_FIELDS + d["fields"]: f[n] = "[%s]" % h
    f.update({"number": fmt_number(d["pattern"], 1, "PRJ", today(), type_), "date": today(), "project": "Sample project", "project_code": "PRJ", "company": co.get("name") or "Sample Company LLC",
              "company_address": co.get("address", ""), "company_phone": co.get("phone", ""), "company_email": co.get("email", ""), "to": "The Main Contractor", "from": "Project Manager",
              "subject": "Sample subject", "status": "Issued", "revision": 0, "generated_on": today(),
              "body": "Dear Sirs,\n\nThis is a sample of the text of the letter or form.\n  - first line\n  - second line\n\nYours faithfully,\nProject Manager"})
    for n in ("gross", "previous_gross", "this_period", "retention", "net", "value", "cost", "rate"): f.setdefault(n, 12345.5)
    for n in ("quantity", "total_quantity"): f[n] = 120.5
    lists = {}
    for ln, cols in d["lists"].items():
        lists[ln] = [{c: ("Sample %s %d" % (c, i + 1) if c not in ("qty", "rate", "mir_qty", "wir_qty", "snag_qty", "pct", "previous", "cumulative", "current", "quantity") else 10.0 * (i + 1)) for c, _ in cols} for i in range(3)]
    return {"fields": f, "lists": lists, "type": type_, "_unknown": set()}

# --------------------------------------------------------------------------------------------
# Rendering dispatcher
# --------------------------------------------------------------------------------------------
def render(type_, ctx, tpl):
    if tpl is None: return "document.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", default_docx(type_, ctx)
    kind = tpl["kind"]; blob = bytes(tpl["blob"]); opts = json.loads(tpl["opts"] or "{}")
    if kind == "docx": return "document.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", fill_docx(blob, ctx)
    if kind == "xlsx": return "document.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", fill_xlsx(blob, ctx)
    if kind == "pdf": return "document.pdf", "application/pdf", fill_pdf(blob, ctx, type_, opts)
    raise ValueError("unknown template kind")

def to_pdf(blob, ext):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe or ext.lower() not in (".docx", ".xlsx"): return None
    try:
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "in" + ext); open(src, "wb").write(blob)
            subprocess.run([exe, "--headless", "-env:UserInstallation=file://" + os.path.join(td, "profile"), "--convert-to", "pdf", "--outdir", td, src],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=90)
            out = os.path.join(td, "in.pdf")
            return open(out, "rb").read() if os.path.exists(out) else None
    except Exception: return None

# --------------------------------------------------------------------------------------------
# Word (.docx): fill a template, or build the default format. Standard library only.
# --------------------------------------------------------------------------------------------
_T = re.compile(r"(<w:t(?:\s[^>]*)?>)([^<]*)(</w:t>)")
_P = re.compile(r"<w:p(?:\s[^>]*)?>.*?</w:p>", re.S)
_TR = re.compile(r"<w:tr(?:\s[^>]*)?>.*?</w:tr>", re.S)
LISTTOK = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\.([A-Za-z0-9_]+)\s*\}\}")

def _t_xml(text): return '<w:t xml:space="preserve">%s</w:t>' % xesc(clean_xml(text))
def _runs(text, rpr=""):
    out = []
    for i, p in enumerate(clean_xml(text).split("\n")):
        if i: out.append("<w:r>%s<w:br/></w:r>" % rpr)
        out.append("<w:r>%s%s</w:r>" % (rpr, _t_xml(p)))
    return "".join(out)

def _val_xml(v):
    """escaped text with line breaks as <w:br/>; used inside an existing <w:t>"""
    parts = clean_xml(v).split("\n")
    return ('</w:t><w:br/><w:t xml:space="preserve">').join(xesc(p) for p in parts)

def _fill_wordxml(xml, ctx):
    # A. paragraphs whose tokens are split over several runs: merge them into one run (keeps the first run's formatting)
    def norm_par(m):
        p = m.group(0); ts = [xunesc(t) for _, t, _ in _T.findall(p)]
        full = "".join(ts)
        if "{{" not in full: return p
        if sum(len(TOK.findall(t)) for t in ts) == len(TOK.findall(full)): return p
        if "<w:drawing" in p or "<w:pict" in p or "<w:fldChar" in p: return p
        ppr = re.search(r"<w:pPr>.*?</w:pPr>", p, re.S); rpr = ""
        for r in re.findall(r"<w:r(?:\s[^>]*)?>.*?</w:r>", p, re.S):
            if "<w:t" in r:
                mm = re.search(r"<w:rPr>.*?</w:rPr>", r, re.S); rpr = mm.group(0) if mm else ""; break
        head = re.match(r"<w:p(?:\s[^>]*)?>", p).group(0)
        return head + (ppr.group(0) if ppr else "") + _runs(full, rpr) + "</w:p>"
    xml = _P.sub(norm_par, xml)
    # B. table rows holding {{list.column}} tokens are repeated once per item
    def rows(m):
        row = m.group(0); txt = "".join(xunesc(t) for _, t, _ in _T.findall(row)); lm = LISTTOK.search(txt)
        if not lm: return row
        ln = _norm(lm.group(1)); items = ctx["lists"].get(ln) or ctx["lists"].get(lm.group(1)) or []
        def one(item):
            def rep(t):
                def r2(x):
                    mm = LISTTOK.fullmatch(x.group(0))
                    if not mm or _norm(mm.group(1)) != ln: return x.group(0)
                    return _val_xml(fmt_val(lookup(ctx, mm.group(1) + "." + mm.group(2), item or {}), ln + "." + _norm(mm.group(2))))
                return TOK.sub(r2, t)
            return _T.sub(lambda mt: _tag_for(mt.group(1), rep(mt.group(2))) + rep(mt.group(2)) + mt.group(3), row)
        return "".join(one(i) for i in items) if items else one(None)
    xml = _TR.sub(rows, xml)
    # C. remaining tokens
    def rep_all(t):
        return TOK.sub(lambda x: _val_xml(fmt_val(lookup(ctx, x.group(1)), _norm(x.group(1)))) if "." not in x.group(1) else "", t)
    return _T.sub(lambda mt: _tag_for(mt.group(1), rep_all(mt.group(2))) + rep_all(mt.group(2)) + mt.group(3), xml)

def _tag_for(tag, text):
    if "xml:space" in tag or not (text.startswith(" ") or text.endswith(" ") or "\n" in text or "</w:t>" in text): return tag
    return '<w:t xml:space="preserve">'

def fill_docx(blob, ctx):
    zin = zipfile.ZipFile(io.BytesIO(blob)); out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if re.match(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$", it.filename):
                data = _fill_wordxml(data.decode("utf-8"), ctx).encode("utf-8")
            zo.writestr(it.filename, data)
    return out.getvalue()

def docx_tokens(blob):
    zin = zipfile.ZipFile(io.BytesIO(blob)); found = []
    for n in zin.namelist():
        if re.match(r"word/(document|header\d*|footer\d*)\.xml$", n):
            xml = zin.read(n).decode("utf-8")
            for p in _P.findall(xml): found += tokens_in("".join(xunesc(t) for _, t, _ in _T.findall(p)))
    return found

def _img_size(path):
    d = open(path, "rb").read(65536)
    if d[:8] == b"\x89PNG\r\n\x1a\n": return struct.unpack(">II", d[16:24]), "png"
    if d[:2] == b"\xff\xd8":
        i = 2
        while i < len(d) - 9:
            if d[i] != 0xFF: i += 1; continue
            mk = d[i + 1]
            if mk in (0xC0, 0xC1, 0xC2): h, w = struct.unpack(">HH", d[i + 5:i + 9]); return (w, h), "jpeg"
            i += 2 + struct.unpack(">H", d[i + 2:i + 4])[0]
    return None, None

def _cell(text, w, bold=False, shade=None, align=None, size=None):
    rpr = "<w:rPr>%s%s</w:rPr>" % ("<w:b/>" if bold else "", '<w:sz w:val="%d"/>' % size if size else "")
    tcpr = '<w:tcW w:w="%d" w:type="dxa"/>' % w + ('<w:shd w:val="clear" w:color="auto" w:fill="%s"/>' % shade if shade else "")
    ppr = '<w:pPr><w:spacing w:before="20" w:after="20"/>%s</w:pPr>' % ('<w:jc w:val="%s"/>' % align if align else "")
    return "<w:tc><w:tcPr>%s</w:tcPr><w:p>%s%s</w:p></w:tc>" % (tcpr, ppr, _runs(text, rpr))

def _table(rows, widths, header=True, borders=True):
    b = ("<w:tblBorders>" + "".join('<w:%s w:val="single" w:sz="4" w:space="0" w:color="999999"/>' % s for s in ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>") if borders else ""
    x = '<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/>%s<w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid>%s</w:tblGrid>' % (sum(widths), b, "".join('<w:gridCol w:w="%d"/>' % w for w in widths))
    for i, r in enumerate(rows):
        x += "<w:tr>" + ("<w:trPr><w:tblHeader/></w:trPr>" if header and i == 0 else "")
        for j, (txt, w) in enumerate(zip(r, widths)):
            num = header and i > 0 and bool(re.fullmatch(r"-?[\d,]+(\.\d+)?", str(txt).strip()))
            x += _cell(txt, w, bold=(header and i == 0), shade="E8EEF0" if header and i == 0 else None, align="right" if num else None, size=18 if len(widths) > 4 else None)
        x += "</w:tr>"
    return x + "</w:tbl>"

def _para(text, bold=False, size=None, space_after=80, rpr_extra=""):
    rpr = "<w:rPr>%s%s%s</w:rPr>" % ("<w:b/>" if bold else "", '<w:sz w:val="%d"/>' % size if size else "", rpr_extra)
    return '<w:p><w:pPr><w:spacing w:before="0" w:after="%d"/></w:pPr>%s</w:p>' % (space_after, _runs(text, rpr))

def default_docx(type_, ctx, tokens=False):
    """The built-in Word format: letterhead, title, reference table, text, and one table per list."""
    f = ctx["fields"]; d = DOC_TYPES[type_]; img = ""; rels = ""; extra = []
    lp = LOGO_PATH() if not tokens else None
    if lp and os.path.exists(lp):
        size, kind = _img_size(lp)
        if size:
            h = 600000; w = int(h * size[0] / size[1]); ext = "png" if kind == "png" else "jpeg"
            extra.append(("word/media/logo." + ext, open(lp, "rb").read()))
            rels = '<Relationship Id="rIdLogo" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/logo.%s"/>' % ext
            img = ('<w:r><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="%d" cy="%d"/><wp:docPr id="1" name="Logo"/>'
                   '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
                   '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:nvPicPr><pic:cNvPr id="0" name="logo"/><pic:cNvPicPr/></pic:nvPicPr>'
                   '<pic:blipFill><a:blip r:embed="rIdLogo"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
                   '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>') % (w, h, w, h)
    v = (lambda k: ("{{%s}}" % k)) if tokens else (lambda k: fmt_val(f.get(k, ""), k))
    body = ""
    if img: body += '<w:p><w:pPr><w:spacing w:after="60"/></w:pPr>%s</w:p>' % img
    body += _para(v("company"), bold=True, size=30, space_after=20)
    body += _para(" | ".join(x for x in ([v("company_address"), v("company_phone"), v("company_email")] if tokens else [f.get("company_address"), f.get("company_phone"), f.get("company_email")]) if x), size=17, space_after=40)
    body += '<w:p><w:pPr><w:pBdr><w:bottom w:val="single" w:sz="12" w:space="1" w:color="0B7A75"/></w:pBdr><w:spacing w:after="160"/></w:pPr></w:p>'
    body += _para((v("subject") if tokens else (f.get("subject") or d["label"])), bold=True, size=28, space_after=120)
    meta = [("Number", v("number")), ("Date", v("date")), ("Project", (v("project") + " (" + v("project_code") + ")") if tokens else "%s (%s)" % (f.get("project", ""), f.get("project_code", ""))),
            ("To", v("to")), ("From", v("from")), ("Related to", v("ref")), ("Status", v("status"))]
    if not tokens: meta = [(a, b) for a, b in meta if str(b).strip() not in ("", "()")]
    body += _table([[a, b] for a, b in meta], [2200, 7160], header=False) + _para("", space_after=120)
    extra_fields = [(n, h) for n, h in d["fields"] if (tokens or f.get(n) not in (None, ""))]
    if extra_fields:
        body += _table([[h, v(n)] for n, h in extra_fields], [3600, 5760], header=False) + _para("", space_after=120)
    body += _para(v("body"), space_after=120, rpr_extra='<w:sz w:val="21"/>')
    for ln, cols in d["lists"].items():
        items = ctx["lists"].get(ln, [])
        if tokens: rows = [[h for _, h in cols], ["{{%s.%s}}" % (ln, c) for c, _ in cols]]
        elif items: rows = [[h for _, h in cols]] + [[fmt_val(it.get(c, ""), ln + "." + c) for c, _ in cols] for it in items]
        else: continue
        wd = int(9360 / len(cols)); body += _para(ln.capitalize(), bold=True, space_after=40) + _table(rows, [wd] * len(cols)) + _para("", space_after=80)
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
           '<w:body>%s<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1000" w:right="1270" w:bottom="1000" w:left="1270" w:header="500" w:footer="500" w:gutter="0"/></w:sectPr></w:body></w:document>' % body)
    styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:docDefaults><w:rPrDefault><w:rPr>'
              '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/><w:sz w:val="21"/><w:lang w:val="en-GB"/></w:rPr></w:rPrDefault></w:docDefaults></w:styles>')
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/><Default Extension="jpeg" ContentType="image/jpeg"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
          '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", doc); z.writestr("word/styles.xml", styles)
        z.writestr("word/_rels/document.xml.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdSt" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>%s</Relationships>' % rels)
        for n, b in extra: z.writestr(n, b)
    return out.getvalue()

# --------------------------------------------------------------------------------------------
# Excel (.xlsx): fill a template (openpyxl), starter template
# --------------------------------------------------------------------------------------------
def _need(mod, pipname=None):
    try: return __import__(mod)
    except ImportError: raise RuntimeError("The Python package '%s' is needed for this format. Run: pip install %s" % (mod, pipname or mod))

def _shift_formula(fm, r, k):
    """rows below the repeated row move down by k; ranges that end on the repeated row grow to include the new rows"""
    def rng(m):
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        n1 = r1 + k if r1 > r else r1; n2 = r2 + k if (r2 > r or r2 == r) else r2
        return "%s%d:%s%d" % (c1, n1, c2, n2)
    fm = re.sub(r"(?<![A-Za-z!'])(\$?[A-Z]{1,3}\$?)(\d+):(\$?[A-Z]{1,3}\$?)(\d+)", lambda m: "\x00" + rng(m).replace(":", "\x01"), fm)
    fm = re.sub(r"(?<![A-Za-z!'\x00\x01\d])(\$?[A-Z]{1,3}\$?)(\d+)(?![\d(])", lambda m: "%s%d" % (m.group(1), int(m.group(2)) + k) if int(m.group(2)) > r else m.group(0), fm)
    return fm.replace("\x00", "").replace("\x01", ":")

def _sub_str(s, ctx, item=None, ln=None):
    def r(m):
        key = m.group(1)
        if "." in key:
            if item is None or _norm(key.split(".")[0]) != ln: return m.group(0) if item is None and ln is None else ""
            return fmt_val(lookup(ctx, key, item), _norm(key))
        return fmt_val(lookup(ctx, key), _norm(key))
    return TOK.sub(r, s)

def _val_for(s, ctx, item=None, ln=None):
    m = TOK.fullmatch(s.strip())
    if m:
        key = m.group(1)
        if "." in key: raw = lookup(ctx, key, item) if (item is not None and _norm(key.split(".")[0]) == ln) else ""
        else: raw = lookup(ctx, key)
        if isinstance(raw, bool): return "Yes" if raw else "No"
        if isinstance(raw, (int, float)): return raw
        return "" if raw is None else str(raw)
    return _sub_str(s, ctx, item, ln)

def fill_xlsx(blob, ctx):
    from copy import copy
    ox = _need("openpyxl"); wb = ox.load_workbook(io.BytesIO(blob))
    for ws in wb.worksheets:
        lrows = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and LISTTOK.search(c.value): lrows.append((c.row, _norm(LISTTOK.search(c.value).group(1)))); break
        for r, ln in sorted(lrows, reverse=True):
            items = ctx["lists"].get(ln) or []; n = max(len(items), 1); k = n - 1
            ncol = ws.max_column; tmpl = [(ws.cell(r, cc).value, copy(ws.cell(r, cc)._style)) for cc in range(1, ncol + 1)]; h = ws.row_dimensions[r].height
            merges = list(ws.merged_cells.ranges); inrow = []; below = []
            for mr in merges:
                if mr.min_row == mr.max_row == r: inrow.append((mr.min_col, mr.max_col))
                elif mr.min_row > r: below.append((mr.min_col, mr.min_row, mr.max_col, mr.max_row))
            if k > 0:
                for mr in merges:
                    if mr.min_row >= r and not (mr.min_row == mr.max_row == r) and mr.min_row > r: ws.unmerge_cells(str(mr))
                for row in ws.iter_rows(min_row=r + 1):
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("="): c.value = _shift_formula(c.value, r, k)
                for rowc in ws.iter_rows(min_row=1, max_row=r):
                    for c in rowc:
                        if isinstance(c.value, str) and c.value.startswith("="): c.value = _shift_formula(c.value, r, k)
                ws.insert_rows(r + 1, k)
                for c1, r1, c2, r2 in below: ws.merge_cells(start_row=r1 + k, start_column=c1, end_row=r2 + k, end_column=c2)
            for i in range(n):
                it = items[i] if items else None
                for cc, (val, st) in enumerate(tmpl, 1):
                    cell = ws.cell(r + i, cc); cell._style = copy(st)
                    if isinstance(val, str):
                        if val.startswith("="):
                            try:
                                from openpyxl.formula.translate import Translator
                                cell.value = Translator(val, origin="A%d" % r).translate_formula("A%d" % (r + i)) if i else val
                            except Exception: cell.value = val
                        else: cell.value = _val_for(val, ctx, it, ln) if LISTTOK.search(val) or TOK.search(val) else val
                    else: cell.value = val
                if h: ws.row_dimensions[r + i].height = h
                if i > 0:
                    for c1, c2 in inrow: ws.merge_cells(start_row=r + i, start_column=c1, end_row=r + i, end_column=c2)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "{{" in c.value: c.value = _val_for(c.value, ctx)
        for hf in (ws.oddHeader, ws.oddFooter, ws.evenHeader, ws.evenFooter, ws.firstHeader, ws.firstFooter):
            for part in (hf.left, hf.center, hf.right):
                if part.text and "{{" in part.text: part.text = _sub_str(part.text, ctx)
    out = io.BytesIO(); wb.save(out); return out.getvalue()

def xlsx_tokens(blob):
    ox = _need("openpyxl"); wb = ox.load_workbook(io.BytesIO(blob)); found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str): found += tokens_in(c.value)
    return found

def starter_xlsx(type_):
    ox = _need("openpyxl"); from openpyxl.styles import Font, PatternFill, Alignment
    wb = ox.Workbook(); ws = wb.active; ws.title = type_; d = DOC_TYPES[type_]
    ws["A1"] = "{{company}}"; ws["A1"].font = Font(bold=True, size=14); ws["A2"] = "{{company_address}}"
    ws["A4"] = "{{subject}}"; ws["A4"].font = Font(bold=True, size=13)
    r = 6
    for lab, key in [("Number", "number"), ("Date", "date"), ("Project", "project"), ("To", "to"), ("From", "from"), ("Related to", "ref"), ("Status", "status")] + [(h, n) for n, h in d["fields"]]:
        ws.cell(r, 1, lab).font = Font(bold=True); ws.cell(r, 2, "{{%s}}" % key); r += 1
    r += 1; ws.cell(r, 1, "{{body}}").alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); ws.row_dimensions[r].height = 120; r += 2
    for ln, cols in d["lists"].items():
        for j, (c, h) in enumerate(cols, 1):
            x = ws.cell(r, j, h); x.font = Font(bold=True); x.fill = PatternFill("solid", fgColor="E8EEF0")
            ws.cell(r + 1, j, "{{%s.%s}}" % (ln, c))
        r += 4
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 34
    for col in "CDEFGHIJK": ws.column_dimensions[col].width = 16
    out = io.BytesIO(); wb.save(out); return out.getvalue()

# --------------------------------------------------------------------------------------------
# PDF: fill form fields, replace {{tokens}} in the page text, or use the PDF as letterhead
# --------------------------------------------------------------------------------------------
def _pdf_modules():
    return _need("pypdf"), _need("reportlab")

def _norm_field(name):
    return _norm(re.sub(r"[\[\]]", "_", str(name))).strip("_")

def _pdf_value(ctx, name):
    n = _norm_field(name); last = _norm_field(re.split(r"[.\\/]", str(name))[-1])
    for cand in (n, last):
        m = re.fullmatch(r"([a-z0-9]+?)[._ ]?([a-z0-9_]+?)[._ #-]*(\d+)", cand)
        f = {_norm(a): b for a, b in ctx["fields"].items()}
        if cand in f: return fmt_val(f[cand], cand)
        for ln, items in ctx["lists"].items():
            mm = re.fullmatch(re.escape(_norm(ln)) + r"[._ -]*([a-z0-9_]+?)[._ #-]+(\d+)", cand)
            if mm:
                i = int(mm.group(2)) - 1; col = mm.group(1)
                return fmt_val(items[i].get(col, ""), col) if 0 <= i < len(items) else ""
    ctx["_unknown"].add(str(name)); return ""

def _pdf_token_boxes(reader):
    boxes = []
    for pi, page in enumerate(reader.pages):
        def visit(text, cm, tm, fd, fs, _pi=pi):
            if text and "{{" in text and TOK.search(text):
                x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]; y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
                size = (fs or 10) * ((tm[2] ** 2 + tm[3] ** 2) ** .5 or 1) * ((cm[2] ** 2 + cm[3] ** 2) ** .5 or 1)
                boxes.append((_pi, x, y, size, text))
        try: page.extract_text(visitor_text=visit)
        except Exception: pass
    return boxes

def fill_pdf(blob, ctx, type_, opts):
    pypdf, rl = _pdf_modules(); from pypdf import PdfReader, PdfWriter
    from reportlab.pdfgen import canvas as rcanvas; from reportlab.pdfbase.pdfmetrics import stringWidth
    reader = PdfReader(io.BytesIO(blob))
    if reader.is_encrypted:
        try: reader.decrypt("")
        except Exception: raise ValueError("The PDF is password protected")
    fields = reader.get_fields() or {}; mode = (opts or {}).get("mode", "auto")
    if fields and mode in ("auto", "form"):
        w = PdfWriter(); w.append(reader)
        vals = {name: _pdf_value(ctx, name) for name in fields}
        for page in w.pages:
            try: w.update_page_form_field_values(page, vals, auto_regenerate=False, flatten=True)
            except TypeError: w.update_page_form_field_values(page, vals, auto_regenerate=False)
        try: w.set_need_appearances_writer(True)
        except Exception: pass
        out = io.BytesIO(); w.write(out); return out.getvalue()
    boxes = _pdf_token_boxes(reader) if mode in ("auto", "tokens") else []
    if boxes:
        w = PdfWriter(); w.append(reader)
        for pi in sorted({b[0] for b in boxes}):
            pg = w.pages[pi]; W, H = float(pg.mediabox.width), float(pg.mediabox.height)
            buf = io.BytesIO(); cv = rcanvas.Canvas(buf, pagesize=(W, H))
            for (p_, x, y, size, text) in [b for b in boxes if b[0] == pi]:
                size = max(6, min(size, 24)); xx = x; plain = ""
                pos = 0
                for m in TOK.finditer(text):
                    pre = text[pos:m.start()]; xx += stringWidth(pre, "Helvetica", size); tw = stringWidth(m.group(0), "Helvetica", size)
                    val = "" if "." in m.group(1) else fmt_val(lookup(ctx, m.group(1)), _norm(m.group(1))); val = val.replace("\n", " ")
                    cv.setFillColorRGB(1, 1, 1); cv.rect(xx - 1, y - size * 0.25, max(tw, stringWidth(val, "Helvetica", size)) + 2, size * 1.2, stroke=0, fill=1)
                    cv.setFillColorRGB(0, 0, 0); cv.setFont("Helvetica", size); cv.drawString(xx, y, val); xx += tw; pos = m.end()
            cv.save(); buf.seek(0); pg.merge_page(PdfReader(buf).pages[0])
        out = io.BytesIO(); w.write(out); return out.getvalue()
    # letterhead: the document text is laid out inside the margins of the uploaded page(s)
    top = float((opts or {}).get("top", 130)); bottom = float((opts or {}).get("bottom", 70))
    content = PdfReader(io.BytesIO(default_pdf(type_, ctx, letterhead=True, top=top, bottom=bottom)))
    w = PdfWriter(); tp = len(reader.pages)
    for i, cp in enumerate(content.pages):
        base = PdfReader(io.BytesIO(blob)).pages[0 if i == 0 or tp < 2 else 1]
        base.merge_page(cp); w.add_page(base)
    out = io.BytesIO(); w.write(out); return out.getvalue()

def pdf_tokens(blob):
    pypdf, _ = _pdf_modules(); from pypdf import PdfReader
    r = PdfReader(io.BytesIO(blob)); fields = list((r.get_fields() or {}).keys())
    if fields: return fields, "form"
    tk = []
    for b in _pdf_token_boxes(r): tk += tokens_in(b[4])
    return tk, ("tokens" if tk else "letterhead")

def default_pdf(type_, ctx, letterhead=False, top=60, bottom=60):
    _, rl = _pdf_modules()
    from reportlab.lib.pagesizes import A4; from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle; from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, HRFlowable
    f = ctx["fields"]; d = DOC_TYPES[type_]; buf = io.BytesIO(); wid = A4[0] - 40 * mm
    sty = ParagraphStyle("b", fontName="Helvetica", fontSize=9.5, leading=12.5); small = ParagraphStyle("s", parent=sty, fontSize=8, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("c", parent=sty, fontSize=8, leading=10)
    def P(t, s=sty):
        t = xesc(clean_xml(t)).replace("\n", "<br/>")
        t = re.sub(r"(?m)(^|<br/>)( +)", lambda m: m.group(1) + "&nbsp;" * len(m.group(2)), t)
        return Paragraph(t, s)
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=top if letterhead else 18 * mm, bottomMargin=bottom if letterhead else 18 * mm)
    st = []
    if not letterhead:
        lp = LOGO_PATH()
        if lp and os.path.exists(lp):
            try:
                size, _k = _img_size(lp); h = 14 * mm; st.append(Image(lp, width=h * size[0] / size[1], height=h, hAlign="LEFT")); st.append(Spacer(1, 3))
            except Exception: pass
        st.append(P(f.get("company", ""), ParagraphStyle("h", parent=sty, fontName="Helvetica-Bold", fontSize=15, leading=18)))
        st.append(P(" | ".join(x for x in (f.get("company_address"), f.get("company_phone"), f.get("company_email")) if x), small))
        st.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0B7A75"), spaceBefore=4, spaceAfter=8))
    st.append(P(f.get("subject") or d["label"], ParagraphStyle("t", parent=sty, fontName="Helvetica-Bold", fontSize=13, leading=16))); st.append(Spacer(1, 6))
    meta = [("Number", f.get("number")), ("Date", f.get("date")), ("Project", "%s (%s)" % (f.get("project", ""), f.get("project_code", ""))), ("To", f.get("to")), ("From", f.get("from")),
            ("Related to", f.get("ref")), ("Status", f.get("status"))] + [(h, fmt_val(f.get(n), n)) for n, h in d["fields"] if f.get(n) not in (None, "")]
    meta = [(a, b) for a, b in meta if str(b or "").strip() not in ("", "()")]
    t = Table([[P(a, cell), P(fmt_val(b), cell)] for a, b in meta], colWidths=[40 * mm, wid - 40 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#999999")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2F3")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    st += [t, Spacer(1, 8), P(f.get("body", ""))]
    for ln, cols in d["lists"].items():
        items = ctx["lists"].get(ln, [])
        if not items: continue
        st += [Spacer(1, 8), P(ln.capitalize(), ParagraphStyle("lh", parent=sty, fontName="Helvetica-Bold"))]
        rows = [[P(h, cell) for _, h in cols]] + [[P(fmt_val(it.get(c, ""), ln + "." + c), cell) for c, _ in cols] for it in items]
        tb = Table(rows, colWidths=[wid / len(cols)] * len(cols), repeatRows=1)
        tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#999999")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF0")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        st.append(tb)
    doc.build(st); return buf.getvalue()

# --------------------------------------------------------------------------------------------
# Template store
# --------------------------------------------------------------------------------------------
def add_template(project, scope_all, type_, name, filename, blob, opts, by):
    if type_ not in DOC_TYPES: raise ValueError("unknown document type")
    if len(blob) > MAX_TEMPLATE: raise ValueError("The file is larger than 8 MB")
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext not in ("docx", "xlsx", "pdf"): raise ValueError("Upload a Word (.docx), Excel (.xlsx) or PDF (.pdf) file. Older .doc/.xls files: save them as .docx/.xlsx first.")
    notes = []
    if ext == "docx":
        try: zipfile.ZipFile(io.BytesIO(blob)).getinfo("word/document.xml")
        except Exception: raise ValueError("This is not a valid Word .docx file")
        toks = docx_tokens(blob); mode = "tokens" if toks else "plain"
    elif ext == "xlsx":
        try: toks = xlsx_tokens(blob)
        except RuntimeError: raise
        except Exception: raise ValueError("This is not a valid Excel .xlsx file")
        mode = "tokens" if toks else "plain"
    else:
        try: toks, mode = pdf_tokens(blob)
        except RuntimeError: raise
        except Exception: raise ValueError("This PDF cannot be read (it may be damaged or password protected)")
    ok = known_tokens(type_)
    if ext == "pdf": unknown = [] if mode != "tokens" else sorted({t for t in toks if _norm(t) not in {_norm(x) for x in ok}})
    else: unknown = sorted({t for t in toks if _norm(t) not in {_norm(x) for x in ok}})
    c = conn()
    try:
        proj = "" if scope_all else project
        c.execute("BEGIN IMMEDIATE")
        c.execute("UPDATE doc_templates SET active=0 WHERE type=? AND project=?", (type_, proj))
        cur = c.execute("INSERT INTO doc_templates (project, type, name, filename, kind, blob, opts, active, uploaded_by, uploaded_at, size) VALUES (?,?,?,?,?,?,?,1,?,?,?)",
                        (proj, type_, (name or filename)[:100], filename[:200], ext, blob, json.dumps(opts or {}), by, now(), len(blob)))
        c.execute("COMMIT"); tid = cur.lastrowid
    except Exception: c.execute("ROLLBACK"); raise
    finally: c.close()
    if ext == "docx" and not toks: notes.append("No {{placeholders}} were found in this Word file. Add tokens such as {{number}} or {{body}}, or download the starter file for this form.")
    if ext == "xlsx" and not toks: notes.append("No {{placeholders}} were found in this Excel file. Add tokens such as {{number}} or {{body}}, or download the starter file for this form.")
    if ext == "pdf": notes.append({"form": "This PDF has form fields: they are filled by field name (for example number, date, to, subject, body).",
                                   "tokens": "{{placeholders}} found in the PDF text are replaced.",
                                   "letterhead": "This PDF has no form fields or {{placeholders}}: it is used as the letterhead and the text of each document is printed inside its margins."}[mode])
    if unknown: notes.append("Not recognised for this form (they will be left blank): " + ", ".join("{{%s}}" % u for u in unknown))
    return {"id": tid, "kind": ext, "mode": mode, "tokens": sorted(set(toks)), "unknown": unknown, "notes": notes}

def list_templates(project):
    c = conn()
    try:
        return [dict(r) for r in c.execute("SELECT id, project, type, name, filename, kind, opts, active, uploaded_by, uploaded_at, size FROM doc_templates WHERE project IN (?, '') ORDER BY type, active DESC, id DESC", (project,))]
    finally: c.close()

def template_row(tid):
    c = conn()
    try: return c.execute("SELECT * FROM doc_templates WHERE id=?", (tid,)).fetchone()
    finally: c.close()

def set_template_active(project, tid, active):
    c = conn()
    try:
        r = c.execute("SELECT * FROM doc_templates WHERE id=? AND project IN (?, '')", (tid, project)).fetchone()
        if not r: raise ValueError("unknown format")
        c.execute("BEGIN IMMEDIATE")
        if active: c.execute("UPDATE doc_templates SET active=0 WHERE type=? AND project=?", (r["type"], r["project"]))
        c.execute("UPDATE doc_templates SET active=? WHERE id=?", (1 if active else 0, tid)); c.execute("COMMIT")
    finally: c.close()

def delete_template(project, tid):
    c = conn()
    try: c.execute("DELETE FROM doc_templates WHERE id=? AND project IN (?, '')", (tid, project))
    finally: c.close()

def test_render(project, tid):
    r = template_row(tid)
    if not r: raise ValueError("unknown format")
    ctx = sample_ctx(r["type"]); fn, mime, blob = render(r["type"], ctx, r)
    return safe_name("sample_" + r["type"]) + os.path.splitext(fn)[1], mime, blob, sorted(ctx["_unknown"])

def starter(type_, kind):
    ctx = sample_ctx(type_)
    if kind == "docx": return "Starter_%s.docx" % type_, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", default_docx(type_, ctx, tokens=True)
    if kind == "xlsx": return "Starter_%s.xlsx" % type_, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", starter_xlsx(type_)
    raise ValueError("starter files are available for Word and Excel")

# --------------------------------------------------------------------------------------------
# Reports
# --------------------------------------------------------------------------------------------
COLS = ["Type", "Number", "Rev", "Date", "Title", "To", "From", "Related to", "Amount / quantity", "Status", "Sent", "Issued by", "Registered", "Note"]
def _report_rows(docs):
    out = []
    for d in docs:
        note = ("VOID: " + (d["void_reason"] or "")) if d["status"] == "Void" else ("Number used twice (created offline)" if d["dup"] else "")
        out.append([d["type"], d["number"], d["rev"], d["date"], d["title"], d["to_party"], d["from_party"], d["ref"], d["amount"], d["status"], "Yes" if d["sent"] else "No", d["created_by"], d["created_at"], note])
    return out

def report(project, fmt, docs, title="Document register", copies="orig"):
    stamp = today(); base = safe_name("%s_%s_%s" % (project, title.replace(" ", "_"), stamp)); rows = _report_rows(docs)
    if fmt == "csv":
        s = io.StringIO(); w = csv.writer(s); w.writerow(COLS); w.writerows(rows)
        return base + ".csv", "text/csv; charset=utf-8", ("﻿" + s.getvalue()).encode("utf-8")
    if fmt == "xlsx": return base + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _register_xlsx(project, rows, docs, title)
    if fmt == "pdf": return base + ".pdf", "application/pdf", _register_pdf(project, rows, title)
    if fmt == "zip":
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("Register.xlsx", _register_xlsx(project, rows, docs, title)) if _has("openpyxl") else z.writestr("Register.csv", report(project, "csv", docs, title)[2])
            used = set()
            for d in docs:
                fl = get_file(project, d["id"], as_="pdf" if copies == "pdf" else "")
                if not fl: continue
                fn, _m, b = fl; name = "%s/%s" % (d["type"], fn); k = 1
                while name in used: k += 1; name = "%s/%s_%d%s" % (d["type"], os.path.splitext(fn)[0], k, os.path.splitext(fn)[1])
                used.add(name); z.writestr(name, b)
        return base + ".zip", "application/zip", out.getvalue()
    raise ValueError("unknown report format")

def _has(m):
    try: __import__(m); return True
    except ImportError: return False

def _register_xlsx(project, rows, docs, title):
    ox = _need("openpyxl"); from openpyxl.styles import Font, PatternFill, Alignment
    wb = ox.Workbook(); ws = wb.active; ws.title = "Register"
    ws.append([title + " - " + PROJECT_NAME(project) + " (" + project + ")"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Generated " + now()]); ws.append([])
    ws.append(COLS)
    for c in ws[4]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8EEF0")
    for r in rows: ws.append(r)
    ws.freeze_panes = "A5"; ws.auto_filter.ref = "A4:%s%d" % (chr(64 + len(COLS)), max(5, len(rows) + 4))
    for i, w in enumerate([8, 16, 5, 11, 40, 24, 20, 24, 18, 12, 6, 16, 16, 30]): ws.column_dimensions[chr(65 + i)].width = w
    s = wb.create_sheet("Summary"); s.append(["Type", "Description", "Total", "Issued / approved", "Draft / pending", "Void"]); [setattr(c, "font", Font(bold=True)) for c in s[1]]
    for t, dd in DOC_TYPES.items():
        sub = [d for d in docs if d["type"] == t]
        if not sub: continue
        void = sum(1 for d in sub if d["status"] == "Void"); done = sum(1 for d in sub if d["status"] in ("Issued", "Approved", "Submitted", "Closed", "Sent"))
        s.append([t, dd["label"], len(sub), done, len(sub) - done - void, void])
    s.column_dimensions["B"].width = 44
    out = io.BytesIO(); wb.save(out); return out.getvalue()

def _register_pdf(project, rows, title):
    _need("reportlab")
    from reportlab.lib.pagesizes import A4, landscape; from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle; from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    buf = io.BytesIO(); cs = ParagraphStyle("c", fontName="Helvetica", fontSize=7, leading=8.5)
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    sel = [0, 1, 3, 4, 5, 8, 9, 10, 12, 13]; wd = [30, 78, 50, 150, 90, 60, 50, 30, 80, 100]
    data = [[Paragraph("<b>%s</b>" % xesc(COLS[i]), cs) for i in sel]] + [[Paragraph(xesc(clean_xml(str(r[i] if r[i] is not None else ""))), cs) for i in sel] for r in rows]
    t = Table(data, colWidths=wd, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .3, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF0")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    hs = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=12)
    doc.build([Paragraph(xesc("%s - %s (%s)" % (title, PROJECT_NAME(project), project)), hs), Paragraph("Generated " + now() + " - %d item(s)" % len(rows), cs), Spacer(1, 8), t])
    return buf.getvalue()
