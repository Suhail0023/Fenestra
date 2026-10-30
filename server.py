#!/usr/bin/env python3
# FenOps - Copyright (c) 2026 Charles Couteau. All rights reserved.
# Proprietary and confidential. Unauthorised copying, modification, distribution or reverse engineering is prohibited. See LICENSE.txt.
"""FenOps local server. One SQLite file (fenestra.db in the data folder) is the database for all projects.
   python server.py            -> http://localhost:8000   (this laptop only)
   python server.py --lan      -> also reachable from phones on the same Wi-Fi
Logins: users.json (passwords stored as salted PBKDF2 hashes). Managed in the app (Users tab, Management role).
Email (optional): SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS, SMTP_FROM.
WhatsApp (optional): WA_TOKEN + WA_PHONE_ID (Meta Cloud API) or TWILIO_SID + TWILIO_TOKEN + TWILIO_WA_FROM.
Deployment: DATA_DIR, PORT, TRUST_PROXY=1, SEED_DEMO_USERS=0, INIT_ADMIN_PASSWORD."""
import datetime as dt, glob, hashlib, json, os, re, secrets, shutil, smtplib, sqlite3, sys, threading, time, base64, uuid
from email.message import EmailMessage
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs
import docs, watch, lib

APP = os.path.dirname(os.path.abspath(__file__)); BASE = os.environ.get("DATA_DIR") or APP; os.makedirs(BASE, exist_ok=True); LOCK = threading.Lock()
USERS_FILE = os.path.join(BASE, "users.json"); COMPANY_FILE = os.path.join(BASE, "company.json"); PROJ_FILE = os.path.join(BASE, "projects.json"); TREE_FILE = os.path.join(BASE, "ptree.json")
KINDS = ["Project", "Sub-project", "Building", "Scope / package", "Phase", "Zone"]
ASSETS = {"/manifest.webmanifest": ("manifest.webmanifest", "application/manifest+json"), "/sw.js": ("sw.js", "text/javascript"),
          "/icon-192.png": ("icon-192.png", "image/png"), "/icon-512.png": ("icon-512.png", "image/png")}

TABLES = {
 "Items": ("items", ["id","p","name","qty","uom","wtg","lm","ac","sc","seq","dur","start","done","st","team","eng","w","rt","t0","face","ea","miss","mailed","code","cr","bc","vg","vo","vr"]),
 "Productivity": ("rates", ["k","q","z"]),
 "Pending_Approvals": ("pend", ["k","i","f","o","n"]),
 "Inspections": ("insp", ["no","tid","n","qty","s","c","d","dt","dd","by","note","ph","sc","cr"]),
 "Daily_Logs": ("logs", ["d","team","t","w","h","q","s","f"]),
 "Obstructions": ("obs", ["d","t","ty","h","n","ph"]),
 "Outbox": ("mail", ["to","n","miss","from","ps","pf","delay","body"]),
 "Employees": ("emps", ["id","name","role","team","site"]),
 "Attendance": ("att", ["d","t","id","name","in","out","g"]),
 "Activity_Log": ("ev", ["d","tm","id","n","ev","by","note","u"]),
 "Chat_Log": ("chat", ["id","d","tm","proj","by","text"]),
 "Snags": ("snags", ["id","no","tid","loc","desc","owner","due","s","d","dc","dr","dw","px","py","ph","by","note"]),
 "Permits": ("ptw", ["id","no","tid","type","vf","vt","s","by","ap","d","note"]),
 "Toolbox_Talks": ("tbt", ["id","d","tid","topic","by","n","note"]),
 "Incidents": ("inc", ["id","d","ty","tid","desc","act","s","ph","by"]),
 "Requests": ("reqs", ["id","to","tid","n","items","d","due","s","reply","dc","ch","b","esc","eb"]),
 "Delays": ("delays", ["id","d","tid","cause","party","days","note","rid"]),
 "Variations": ("vars", ["id","no","d","desc","bid","dq","val","s","note","wir"]),
 "Costs": ("costs", ["id","d","bid","type","amt","note"]),
 "Drawings": ("draws", ["id","name","url"])}
CONFIG_KEYS = ["proj", "base", "baseH", "xl", "wirTpl", "contacts", "fin", "diary", "pp", "pts", "awd", "lean", "fid", "sp", "dp", "mh", "vo", "ipa", "ipas", "mt", "tl", "sir", "wp", "cfg", "sigs", "roll", "fca", "la", "wh", "warn", "hrStart", "tc", "dg", "lgt", "rw", "rwStart", "dlg", "geo", "kpi", "kpir", "kpic", "onb", "eng"]
INTS = ("id","p","seq","dur","i","t","tid","dw","rid","bid")
LISTF = {"insp": ("c", "cr", "ph"), "obs": ("ph",), "snags": ("ph",), "inc": ("ph",), "reqs": ("items",)}
DB_FILE = os.path.join(BASE, "fenestra.db")

def code_of(p):
    c = re.sub(r"[^A-Za-z0-9_-]", "", p or "") or "LR001"
    return c[:24]
def load_json(fp, default):
    try: return json.load(open(fp, encoding="utf-8"))
    except Exception: return default
def save_json(fp, obj):
    tmp = fp + ".tmp"; json.dump(obj, open(tmp, "w", encoding="utf-8"), indent=1); os.replace(tmp, fp)

# ---------------- SQLite storage ----------------
def q_(name): return '"%s"' % name.replace('"', '""')
def connect():
    c = sqlite3.connect(DB_FILE, timeout=30)
    c.execute("PRAGMA journal_mode=WAL"); c.execute("PRAGMA synchronous=FULL")
    return c
def init_db():
    c = connect()
    with c:
        c.execute("CREATE TABLE IF NOT EXISTS projects (code TEXT PRIMARY KEY, name TEXT, parent TEXT DEFAULT '', kind TEXT DEFAULT 'Project')")
        c.execute("CREATE TABLE IF NOT EXISTS settings (project TEXT NOT NULL, key TEXT NOT NULL, value, PRIMARY KEY (project, key))")
        c.execute("CREATE TABLE IF NOT EXISTS config (project TEXT NOT NULL, key TEXT NOT NULL, value TEXT, PRIMARY KEY (project, key))")
        c.execute("CREATE TABLE IF NOT EXISTS checklist (project TEXT NOT NULL, task_id TEXT NOT NULL, done TEXT, PRIMARY KEY (project, task_id))")
        c.execute("CREATE TABLE IF NOT EXISTS flags_ack (project TEXT NOT NULL, flag_key TEXT NOT NULL, done INTEGER, PRIMARY KEY (project, flag_key))")
        c.execute("CREATE TABLE IF NOT EXISTS todo (id INTEGER PRIMARY KEY, project TEXT NOT NULL, text TEXT)")
        c.execute("CREATE INDEX IF NOT EXISTS ix_todo_project ON todo(project)")
        for sh, (key, cols) in TABLES.items():
            t = sh.lower()
            # no declared column types: SQLite keeps each value as given (int, real or text)
            c.execute("CREATE TABLE IF NOT EXISTS %s (_rid INTEGER PRIMARY KEY, project TEXT NOT NULL, _meta TEXT, %s)" % (q_(t), ", ".join(q_(x) for x in cols)))
            c.execute("CREATE INDEX IF NOT EXISTS %s ON %s(project)" % (q_("ix_%s_project" % t), q_(t)))
            have = {r[1] for r in c.execute("PRAGMA table_info(%s)" % q_(t))}
            for x in cols:   # a column added in a newer version is added to an existing database
                if x not in have: c.execute("ALTER TABLE %s ADD COLUMN %s" % (q_(t), q_(x)))
    c.close()

def project_name(code):
    c = connect()
    try:
        r = c.execute("SELECT name FROM projects WHERE code=?", (code_of(code),)).fetchone()
        return (r[0] if r and r[0] else code)
    finally: c.close()
def logo_path():
    for f in glob.glob(os.path.join(BASE, "logo.*")):
        if f.lower().endswith((".png", ".jpg", ".jpeg")): return f
    return None

def project_exists(p):
    c = connect()
    try: return c.execute("SELECT 1 FROM projects WHERE code=?", (code_of(p),)).fetchone() is not None
    finally: c.close()

def read_state(p="LR001"):
    p = code_of(p); c = connect()
    try:
        S = {}
        for sh, (key, cols) in TABLES.items():
            out = []
            for r in c.execute("SELECT _meta, %s FROM %s WHERE project=? ORDER BY _rid" % (", ".join(q_(x) for x in cols), q_(sh.lower())), (p,)):
                meta = json.loads(r[0]) if r[0] else {}; row = {}
                for col, v in zip(cols, r[1:]):
                    if v in (None, ""): continue
                    if col in meta.get("j", ()): v = json.loads(v)
                    elif col in meta.get("b", ()): v = bool(v)
                    row[col] = v
                if row: out.append(row)
            S[key] = out
        for key, fs in LISTF.items():
            for x in S[key]:
                for f in fs: x[f] = x[f] if isinstance(x.get(f), list) else []
        st = {k: v for k, v in c.execute("SELECT key, value FROM settings WHERE project=?", (p,))}
        ids = [i["id"] for i in S["items"] if "id" in i]
        S["n"] = max(int(st.get("n") or 1), (max(ids) + 1) if ids else 1)
        for key in ("snags", "ptw", "tbt", "inc", "reqs", "delays", "vars", "costs", "draws"):
            for x in S[key]:
                if isinstance(x.get("id"), int): S["n"] = max(S["n"], x["id"] + 1)
        S["ratio"] = {k: int(st.get(k) or d) for k, d in (("sup", 8), ("log", 12), ("store", 20))}
        if st.get("chkitems"): S["chkItems"] = json.loads(st["chkitems"])
        S["teams"] = json.loads(st["teams"]) if st.get("teams") else []
        S["chk"] = {str(k): json.loads(v or "[]") for k, v in c.execute("SELECT task_id, done FROM checklist WHERE project=?", (p,))}
        S["ack"] = {str(k): bool(v) for k, v in c.execute("SELECT flag_key, done FROM flags_ack WHERE project=?", (p,))}
        S["todo"] = [r[0] for r in c.execute("SELECT text FROM todo WHERE project=? ORDER BY id", (p,))]
        for k, v in c.execute("SELECT key, value FROM config WHERE project=?", (p,)):
            if k in CONFIG_KEYS:
                try: S[k] = json.loads(v)
                except Exception: pass
        return S
    finally: c.close()

def _enc(v, col, meta):
    if isinstance(v, bool): meta["b"].append(col); return int(v)
    if isinstance(v, (list, dict, tuple)):
        if not v: return None
        meta["j"].append(col); return json.dumps(v, ensure_ascii=False)
    return v

def write_state(S, p="LR001"):
    p = code_of(p); c = connect()
    try:
        with c:   # one transaction: all of the project's data is replaced together or not at all
            nm = (S.get("proj") or {}).get("name")
            c.execute("INSERT OR IGNORE INTO projects (code, name) VALUES (?, ?)", (p, nm or p))
            if nm: c.execute("UPDATE projects SET name=? WHERE code=? AND IFNULL(name,'')<>?", (nm, p, nm))
            for sh, (key, cols) in TABLES.items():
                t = sh.lower(); c.execute("DELETE FROM %s WHERE project=?" % q_(t), (p,))
                ins = "INSERT INTO %s (project, _meta, %s) VALUES (?, ?, %s)" % (q_(t), ", ".join(q_(x) for x in cols), ", ".join("?" * len(cols)))
                for x in S.get(key, []):
                    meta = {"b": [], "j": []}; vals = [_enc(x.get(col), col, meta) for col in cols]
                    c.execute(ins, [p, json.dumps({k: v for k, v in meta.items() if v}) if (meta["b"] or meta["j"]) else None] + vals)
            c.execute("DELETE FROM settings WHERE project=?", (p,))
            sets = [("n", S.get("n", 1))] + list(S.get("ratio", {}).items()) + [("teams", json.dumps(S.get("teams", []), ensure_ascii=False))]
            if S.get("chkItems"): sets.append(("chkitems", json.dumps(S["chkItems"], ensure_ascii=False)))
            c.executemany("INSERT INTO settings (project, key, value) VALUES (?, ?, ?)", [(p, k, v) for k, v in sets])
            c.execute("DELETE FROM checklist WHERE project=?", (p,))
            c.executemany("INSERT INTO checklist (project, task_id, done) VALUES (?, ?, ?)", [(p, str(k), json.dumps([bool(x) for x in v])) for k, v in S.get("chk", {}).items()])
            c.execute("DELETE FROM flags_ack WHERE project=?", (p,))
            c.executemany("INSERT INTO flags_ack (project, flag_key, done) VALUES (?, ?, 1)", [(p, str(k)) for k, v in S.get("ack", {}).items() if v])
            c.execute("DELETE FROM todo WHERE project=?", (p,))
            c.executemany("INSERT INTO todo (project, text) VALUES (?, ?)", [(p, t_) for t_ in S.get("todo", [])])
            c.execute("DELETE FROM config WHERE project=?", (p,))
            c.executemany("INSERT INTO config (project, key, value) VALUES (?, ?, ?)", [(p, k, json.dumps(S[k], ensure_ascii=False)) for k in CONFIG_KEYS if S.get(k) is not None])
    finally: c.close()

def get_tree():
    c = connect()
    try: return {code: {"parent": par or "", "kind": kind or "Project"} for code, par, kind in c.execute("SELECT code, parent, kind FROM projects")}
    finally: c.close()
def set_tree(code, parent, kind):
    c = connect()
    try:
        with c: c.execute("UPDATE projects SET parent=?, kind=? WHERE code=?", (parent or "", kind, code))
    finally: c.close()

# ---------------- delete a project (archived first) ----------------
TOMB_FILE = os.path.join(BASE, "deleted_projects.json")
def tombstoned(code): return code_of(code) in load_json(TOMB_FILE, [])
def _tomb(code, on):
    t = [x for x in load_json(TOMB_FILE, []) if x != code]
    if on: t.append(code)
    save_json(TOMB_FILE, t)
def build_project_zip(code, who="", reason=""):
    """Zip with the project's state, library files, register files, table rows and engineering records. Returns bytes."""
    import io, zipfile
    code = code_of(code); c = connect(); bio = io.BytesIO()
    try:
        tabs = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")]
        pt = [t for t in tabs if t != "projects" and any(r[1] == "project" for r in c.execute("PRAGMA table_info(%s)" % q_(t)))]
        with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("state.json", json.dumps(read_state(code), ensure_ascii=False, indent=1))
            z.writestr("project.json", json.dumps({"code": code, "name": project_name(code), "by": who, "at": time.strftime("%Y-%m-%d %H:%M:%S"), "reason": reason}, indent=1))
            rows = {}
            for t in pt:
                cur = c.execute("SELECT * FROM %s WHERE project=?" % q_(t), (code,)); names = [x[0] for x in cur.description]
                rows[t] = [{k: ("<%d bytes>" % len(v) if isinstance(v, (bytes, bytearray)) else v) for k, v in zip(names, r)} for r in cur.fetchall()]
            z.writestr("tables.json", json.dumps(rows, ensure_ascii=False, default=str))
            if "library" in tabs:
                for i, fnm, blob in c.execute("SELECT id, filename, data FROM library WHERE project=? AND data IS NOT NULL", (code,)).fetchall():
                    z.writestr("library/%s_%s" % (i, os.path.basename(fnm or "file")), blob)
            if "doc_files" in tabs and "documents" in tabs:
                for i, fnm, blob in c.execute("SELECT f.doc_id, f.filename, f.blob FROM doc_files f JOIN documents d ON d.id=f.doc_id WHERE d.project=? AND f.blob IS NOT NULL", (code,)).fetchall():
                    z.writestr("documents/%s_%s" % (i, os.path.basename(fnm or "file")), blob)
            ef = engrec_path(code)
            if os.path.exists(ef): z.write(ef, "eng_records.jsonl")
    finally: c.close()
    return bio.getvalue()
def save_archive(code, folder, who="", reason=""):
    code = code_of(code); d = os.path.join(BASE, folder); os.makedirs(d, exist_ok=True)
    fn = "%s_%s.zip" % (code, time.strftime("%Y%m%d-%H%M%S")); data = build_project_zip(code, who, reason)
    open(os.path.join(d, fn), "wb").write(data); return fn, data
def delete_project(code, who=""):
    """Archives the project into deleted/<code>_<time>.zip, then removes it. Returns the archive file name."""
    code = code_of(code); fn, _ = save_archive(code, "deleted", who, "project deleted"); c = connect()
    try:
        tabs = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")]
        pt = [t for t in tabs if t != "projects" and any(r[1] == "project" for r in c.execute("PRAGMA table_info(%s)" % q_(t)))]
        with c:
            if "doc_files" in tabs and "documents" in tabs:
                c.execute("DELETE FROM doc_files WHERE doc_id IN (SELECT id FROM documents WHERE project=?)", (code,))
            for t in pt: c.execute("DELETE FROM %s WHERE project=?" % q_(t), (code,))
            c.execute("DELETE FROM projects WHERE code=?", (code,))
    finally: c.close()
    try:
        ef = engrec_path(code)
        if os.path.exists(ef): os.remove(ef)
    except OSError: pass
    _tomb(code, True)
    return fn

def list_projects():
    c = connect()
    try: rows_ = c.execute("SELECT code, name, parent, kind FROM projects ORDER BY code").fetchall()
    finally: c.close()
    codes = {r[0] for r in rows_}; out = []
    for code, name, par, kind in rows_:
        out.append({"code": code, "name": name or code, "parent": par if par in codes and par != code else "", "kind": kind if kind in KINDS else "Project"})
    return out or [{"code": "LR001", "name": "LR001", "parent": "", "kind": "Project"}]

# ---------------- one-time import of legacy Excel databases ----------------
_XL_CLEAN_INTS = INTS
def _xl_clean(k, v):
    if isinstance(v, dt.datetime) or isinstance(v, dt.date): return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time): return v.strftime("%H:%M")
    if isinstance(v, float) and v == int(v) and k in _XL_CLEAN_INTS: return int(v)
    return v.strip() if isinstance(v, str) else v
def _xl_split(s, sep): return [t.strip() for t in str(s).split(sep) if t.strip()]
def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True); S = {}
    def rows(ws): return [r for r in ws.iter_rows(min_row=2, values_only=True) if any(x not in (None, "") for x in r)]
    for sh, (key, cols) in TABLES.items():
        S[key] = [{c: _xl_clean(c, v) for c, v in zip(cols, r) if v not in (None, "")} for r in rows(wb[sh])] if sh in wb.sheetnames else []
    for x in S["items"]:
        if "ea" in x: x["ea"] = [t.strip() == "1" for t in str(x["ea"]).split(";")]
        if "miss" in x: x["miss"] = _xl_split(x["miss"], "|")
    for x in S["mail"]:
        if "miss" in x: x["miss"] = _xl_split(x["miss"], "|")
    for key, fs in LISTF.items():
        for x in S[key]:
            for f in fs: x[f] = _xl_split(x.get(f, ""), ";")
    st = {r[0]: r[1] for r in rows(wb["Settings"])} if "Settings" in wb.sheetnames else {}
    S["n"] = int(st.get("n") or 1)
    S["ratio"] = {k: int(st.get(k) or d) for k, d in (("sup", 8), ("log", 12), ("store", 20))}
    if st.get("chkitems"): S["chkItems"] = _xl_split(st["chkitems"], "|")
    S["teams"] = _xl_split(st.get("teams", ""), ",")
    S["chk"] = {str(r[0]): [x == "1" for x in str(r[1] or "").split(";") if x != ""] for r in rows(wb["Checklist"])} if "Checklist" in wb.sheetnames else {}
    S["ack"] = {str(r[0]): bool(r[1]) for r in rows(wb["Flags_Ack"])} if "Flags_Ack" in wb.sheetnames else {}
    S["todo"] = [r[0] for r in rows(wb["Todo"])] if "Todo" in wb.sheetnames else []
    if "Config" in wb.sheetnames:
        parts = {}
        for k, v in ((r[0], r[1]) for r in rows(wb["Config"])):
            b_, _, n_ = str(k).partition("#"); parts.setdefault(b_, {})[int(n_ or 0)] = v or ""
        for k, pp in parts.items():
            if k in CONFIG_KEYS:
                try: S[k] = json.loads("".join(pp[i] for i in sorted(pp)))
                except Exception: pass
    return S

def migrate_legacy():
    """Copy every AZAL_<CODE>_Project_DB.xlsx that is not yet in the SQL database, plus projects.json / ptree.json.
    The Excel files are left untouched (rename or delete them once you have checked the app)."""
    files = sorted(glob.glob(os.path.join(BASE, "AZAL_*_Project_DB.xlsx")))
    todo = []
    for f in files:
        m = re.match(r"AZAL_(.+)_Project_DB\.xlsx$", os.path.basename(f))
        if m and not project_exists(m.group(1)): todo.append((m.group(1), f))
    if not todo: return
    try: import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("Found Excel project files (%s) that are not yet in fenestra.db.\nRun once:  pip install openpyxl   then start the server again to import them.\n(Not starting, so that your data is not hidden behind an empty project.)" % ", ".join(os.path.basename(f) for _, f in todo))
    names = load_json(PROJ_FILE, {}); tree = load_json(TREE_FILE, {})
    for code, f in todo:
        S = read_xlsx(f)
        if names.get(code): S.setdefault("proj", {}); S["proj"] = S["proj"] if S["proj"].get("name") else dict(S["proj"], name=names[code])
        write_state(S, code); print("Imported %s into fenestra.db" % os.path.basename(f))
    for code, _ in todo:
        t = tree.get(code) or {}
        if t.get("parent") and project_exists(t["parent"]) and t["parent"] != code:
            set_tree(code, t["parent"], t.get("kind") if t.get("kind") in KINDS else "Project")
        elif t.get("kind") in KINDS: set_tree(code, "", t["kind"])

# ---------------- users and tokens ----------------
ROLE_DEFAULTS = [("management", "Management", "Management"), ("pm", "Project manager", "Project manager"), ("planner", "Planning engineer", "Planning engineer"),
                 ("siteeng", "Site engineer", "Site engineer"), ("siteTL", "Site team leader", "Site team leader"), ("qe", "QA/QC", "QA/QC"), ("hse", "HSE", "HSE"),
                 ("logistics", "Site logistics", "Site logistics"),
                 ("contractor", "Main contractor", "Main contractor"), ("production", "Production manager", "Production manager"), ("prodeng", "Production engineer", "Production engineer"), ("foreman", "Production foreman", "Production foreman"), ("prodtl", "Production team leader", "Production team leader"), ("maint", "Maintenance in-charge", "Maintenance in-charge"),
                 ("engman", "Engineering manager", "Engineering manager"), ("teamlead", "Engineering team leader", "Engineering team leader"), ("designer", "Designer", "Designer"),
                 ("procurement", "Procurement manager", "Procurement manager"), ("procofficer", "Procurement officer", "Procurement officer")]
TOKENS = {}
def hash_pw(pw, salt=None):
    salt = salt or secrets.token_hex(8)
    return salt, hashlib.pbkdf2_hmac("sha256", pw.encode(), salt.encode(), 120000).hex()
def users():
    u = load_json(USERS_FILE, None)
    if u is None:
        u = []
        demo = os.environ.get("SEED_DEMO_USERS", "1") != "0"
        for un, name, role in (ROLE_DEFAULTS if demo else ROLE_DEFAULTS[:1]):
            pw = "azal2026" if demo else (os.environ.get("INIT_ADMIN_PASSWORD") or secrets.token_urlsafe(9))
            if not demo: print("Management account created: user name 'management', password:", pw, "(change it at first login)")
            s, h = hash_pw(pw); u.append({"u": un, "name": name, "role": role, "email": "", "phone": "", "salt": s, "hash": h, "must": True})
        save_json(USERS_FILE, u)
    # role names changed in v23: keep existing logins working
    REN = {"Team leader": "Engineering team leader", "Design engineer": "Designer", "Quality engineer": "QA/QC", "HSE officer": "HSE", "Production department": "Production manager", "Production": "Production manager", "Junior designer": "Designer", "Document controller": "Engineering manager"}; ch = False
    for x in u:
        if x.get("role") in REN: x["role"] = REN[x["role"]]; ch = True
    if ch: save_json(USERS_FILE, u)
    # permanent administrator account (Management role), created once; the password is stored only as a salted hash
    mk = os.path.join(BASE, ".admin_seeded")
    if not os.path.exists(mk) and not any(x["u"].lower() == "couteau" for x in u):
        u.append({"u": "Couteau", "name": "Couteau", "role": "Management", "email": "charles.couteau@gmail.com", "phone": "",
                  "salt": "959b64febf8da9d5", "hash": "0b8f9a04807fd85321a752f9856cbddec56005b45893daa9e42b69aff9ba4de1", "must": False})
        save_json(USERS_FILE, u)
        try: open(mk, "w").write("1")
        except OSError: pass
    # super administrator (Management role = every module, user admin, delete project): created once from SUPERADMIN_USER / SUPERADMIN_PASSWORD.
    sp = os.environ.get("SUPERADMIN_PASSWORD", "")
    if sp and len(sp) >= 8:
        sn = (os.environ.get("SUPERADMIN_USER") or "superadmin").strip()
        if not any(x["u"].lower() == sn.lower() for x in u):
            sa, h = hash_pw(sp); u.append({"u": sn, "name": "Super admin", "role": "Management", "email": "", "phone": "", "salt": sa, "hash": h, "must": False, "super": True})
            save_json(USERS_FILE, u)
    return u
def pub(u): return {k: u.get(k, "") for k in ("u", "name", "role", "email", "phone")} | {"must": bool(u.get("must")), "status": u.get("status", "active")}
KNOWN_ROLES = [r for _, _, r in ROLE_DEFAULTS]
LOGO_EXT = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp"}
CO_FIELDS = ("name", "address", "phone", "email", "web", "trn", "footer", "chan")
def company(): return load_json(COMPANY_FILE, {})
SIGNUPS = []
FAILS = {}
ASKS = {}
def audit_tail(n=500):
    """Last n audit lines, newest first: [{at, ev, who, ip}]"""
    fp = os.path.join(BASE, "audit.log"); out = []
    try:
        with open(fp, encoding="utf-8", errors="replace") as f: lines = f.readlines()[-max(1, min(n, 5000)):]
    except OSError: return out
    for ln in reversed(lines):
        p = ln.rstrip("\n").split("\t")
        if len(p) >= 2: out.append({"at": p[0], "ev": p[1], "who": p[2] if len(p) > 2 else "", "ip": p[3] if len(p) > 3 else ""})
    return out


def audit(ev, who="", ip=""):
    try:
        with open(os.path.join(BASE, "audit.log"), "a", encoding="utf-8") as f: f.write("%s\t%s\t%s\t%s\n" % (time.strftime("%Y-%m-%d %H:%M:%S"), ev, who, ip))
    except Exception: pass
def locked(key, lim=5):
    now = time.time(); FAILS[key] = [t for t in FAILS.get(key, []) if now - t < 600]
    return len(FAILS[key]) >= lim and int(600 - (now - FAILS[key][0])) + 1
def user_by_name_or_role(n):
    for u in users():
        if (u.get("email") or u.get("phone")) and (u["name"] == n or u["role"] == n): return u
    return None

def smtp_send(to, subject, body):
    sink = os.environ.get("SMTP_SINK")   # testing aid: append messages to this file instead of sending them
    if sink:
        with open(sink, "a", encoding="utf-8") as f: f.write(json.dumps({"to": to, "subject": subject, "body": body}, ensure_ascii=False) + "\n")
        return
    host = os.environ.get("SMTP_HOST")
    if not host: raise RuntimeError("Email is not configured on this server (set SMTP_HOST, SMTP_USER, SMTP_PASS).")
    e = EmailMessage(); e["From"] = os.environ.get("SMTP_FROM", os.environ.get("SMTP_USER", "")); e["To"] = to; e["Subject"] = subject; e.set_content(body)
    port = int(os.environ.get("SMTP_PORT", 587))
    if port == 465:
        with smtplib.SMTP_SSL(host, port, timeout=25) as s:
            if os.environ.get("SMTP_USER"): s.login(os.environ["SMTP_USER"], os.environ.get("SMTP_PASS", ""))
            s.send_message(e)
    else:
        with smtplib.SMTP(host, port, timeout=25) as s:
            s.starttls()
            if os.environ.get("SMTP_USER"): s.login(os.environ["SMTP_USER"], os.environ.get("SMTP_PASS", ""))
            s.send_message(e)

def wa_provider():
    if os.environ.get("WA_TOKEN") and os.environ.get("WA_PHONE_ID"): return "meta"
    if os.environ.get("TWILIO_SID") and os.environ.get("TWILIO_TOKEN") and os.environ.get("TWILIO_WA_FROM"): return "twilio"
    return ""
def wa_send(phone, text):
    import urllib.request as ur, urllib.parse as up, urllib.error as ue
    pv = wa_provider(); ph = re.sub(r"\D", "", str(phone or ""))
    if not pv: raise RuntimeError("WhatsApp is not configured on this server (set WA_TOKEN and WA_PHONE_ID, or the TWILIO_ variables).")
    if not 8 <= len(ph) <= 15: raise RuntimeError("The WhatsApp number is not valid. Use the international format, for example 9715xxxxxxxx.")
    text = str(text)[:3500]
    try:
        if pv == "meta":
            tpl = os.environ.get("WA_TEMPLATE")
            if tpl:
                msg = {"messaging_product": "whatsapp", "to": ph, "type": "template", "template": {"name": tpl, "language": {"code": os.environ.get("WA_TEMPLATE_LANG", "en")}, "components": [{"type": "body", "parameters": [{"type": "text", "text": re.sub(r"\s*\n+\s*", " | ", text)[:1000]}]}]}}
            else: msg = {"messaging_product": "whatsapp", "to": ph, "type": "text", "text": {"body": text}}
            r = ur.Request("https://graph.facebook.com/v20.0/%s/messages" % os.environ["WA_PHONE_ID"], data=json.dumps(msg).encode(), headers={"Authorization": "Bearer " + os.environ["WA_TOKEN"], "Content-Type": "application/json"})
        else:
            fr = os.environ["TWILIO_WA_FROM"]; fr = fr if fr.startswith("whatsapp:") else "whatsapp:" + fr
            data = up.urlencode({"From": fr, "To": "whatsapp:+" + ph, "Body": text}).encode()
            tok = base64.b64encode((os.environ["TWILIO_SID"] + ":" + os.environ["TWILIO_TOKEN"]).encode()).decode()
            r = ur.Request("https://api.twilio.com/2010-04-01/Accounts/%s/Messages.json" % os.environ["TWILIO_SID"], data=data, headers={"Authorization": "Basic " + tok})
        ur.urlopen(r, timeout=25).read()
    except ue.HTTPError as x:
        raise RuntimeError("WhatsApp provider refused the message (%s): %s" % (x.code, x.read()[:200].decode("utf-8", "replace")))

def deliver(email, phone, subject, body, channel=None):
    """Send by email, WhatsApp or both. Returns {channel: 'sent' | 'skipped: reason' | 'failed: reason'}."""
    ch = channel or company().get("chan") or "email"; res = {}
    if ch in ("email", "both"):
        if not email: res["email"] = "skipped: no email address"
        elif not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email): res["email"] = "failed: invalid email address"
        else:
            try: smtp_send(email, subject, body); res["email"] = "sent"
            except Exception as x: res["email"] = "failed: %s" % x
    if ch in ("whatsapp", "both"):
        if not phone: res["whatsapp"] = "skipped: no WhatsApp number"
        else:
            try: wa_send(phone, (subject + "\n\n" if subject else "") + body); res["whatsapp"] = "sent"
            except Exception as x: res["whatsapp"] = "failed: %s" % x
    return res
SENDS = {}

def notify(mails):
    if not (os.environ.get("SMTP_HOST") or wa_provider()) or not mails: return
    threading.Thread(target=_notify, args=(list(mails),), daemon=True).start()
def _notify(mails):
    for m in mails:
        if m.get("miss") or m.get("body"): continue
        u = user_by_name_or_role(m.get("to"))
        if not u: continue
        r = deliver(u.get("email"), u.get("phone"), "Please accept site task: " + m["n"], "A task has been delegated to you: " + m["n"] + "\nOpen FenOps and press Accept.")
        if any(v.startswith("failed") for v in r.values()): print("notification failed:", r)

COPYRIGHT_HOLDER = os.environ.get("COPYRIGHT_HOLDER", "Charles Couteau")

# ---- append-only engineering records (instructions, site dimensions, notifications): one JSON line per record, hash-chained
def engrec_path(p): return os.path.join(BASE, "eng_records_%s.jsonl" % re.sub(r"[^A-Za-z0-9_-]", "", p or "LR001"))
def engrec_all(p):
    f = engrec_path(p)
    if not os.path.exists(f): return []
    out = []
    for ln in open(f, encoding="utf-8"):
        ln = ln.strip()
        if ln:
            try: out.append(json.loads(ln))
            except ValueError: pass
    return out
def engrec_add(p, rec, who, role):
    clean = {k: rec.get(k) for k in ("kind", "to", "subject", "body", "sid", "x")}
    clean["subject"] = str(clean.get("subject") or "")[:300]; clean["body"] = str(clean.get("body") or "")[:8000]; clean["kind"] = str(clean.get("kind") or "Notification")[:60]
    with LOCK:
        allr = engrec_all(p); prev = allr[-1] if allr else {}
        r = {"n": (prev.get("n", 0) + 1), "t": dt.datetime.now().astimezone().isoformat(timespec="seconds"), "by": who, "role": role, **clean, "prev": prev.get("h", "")}
        r["h"] = hashlib.sha256(json.dumps({k: v for k, v in r.items() if k != "h"}, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        with open(engrec_path(p), "a", encoding="utf-8") as fh: fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    return r
def engrec_verify(p):
    prev = ""; 
    for r in engrec_all(p):
        h = hashlib.sha256(json.dumps({k: v for k, v in r.items() if k != "h"}, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        if r.get("prev") != prev or r.get("h") != h: return {"ok": False, "broken_at": r.get("n")}
        prev = r["h"]
    return {"ok": True, "count": len(engrec_all(p))}

PRICE_ROLES = {"Management", "Project manager"}   # only these roles receive or can change BOQ prices
PRICE_KEYS = ("cr", "bc", "vr")
def strip_prices(St):
    for it in St.get("items") or []:
        for k in PRICE_KEYS: it.pop(k, None)
    return St
def keep_prices(new, old):
    o = {i.get("id"): i for i in (old.get("items") or [])}
    for it in new.get("items") or []:
        for k in PRICE_KEYS: it.pop(k, None)
        src = o.get(it.get("id"))
        if src:
            for k in PRICE_KEYS:
                if k in src: it[k] = src[k]
    return new
NO_KPI_ROLES = {"Main contractor"}      # roles that are never sent the KPI / certificate module
_APP_CACHE = {}
def app_payload(role):
    """The application page, assembled per role after login. Without a dist/ folder (development) the readable index.html is sent."""
    d = os.path.join(APP, "dist")
    if not os.path.exists(os.path.join(d, "app.html")):
        return open(os.path.join(APP, "index.html"), "rb").read()
    def part(n):
        if n not in _APP_CACHE: _APP_CACHE[n] = open(os.path.join(d, n), encoding="utf-8").read()
        return _APP_CACHE[n]
    scripts = "<script>window.__CR=%s</script>" % json.dumps(COPYRIGHT_HOLDER).replace("</", "<\\/")
    scripts += "<script>" + part("a.js") + "</script>"
    if role not in NO_KPI_ROLES: scripts += "<script>" + part("k.js") + "</script>"
    scripts += "<script>" + part("b.js") + "</script>"
    return part("app.html").replace("<!--#SCRIPTS-->", scripts).encode("utf-8")

class H(BaseHTTPRequestHandler):
    def out(s, code, body, ct="application/json", extra=None):
        b = body if isinstance(body, bytes) else body.encode(); s.send_response(code); s.send_header("Content-Type", ct)
        for k_, v_ in (extra or {}).items(): s.send_header(k_, v_)
        s.send_header("Content-Length", str(len(b))); s.send_header("X-Content-Type-Options", "nosniff"); s.send_header("X-Frame-Options", "SAMEORIGIN"); s.send_header("Referrer-Policy", "same-origin")
        if s.headers.get("X-Forwarded-Proto") == "https": s.send_header("Strict-Transport-Security", "max-age=31536000")
        s.end_headers(); s.wfile.write(b)
    def me(s):
        t = s.headers.get("X-Token", ""); k = TOKENS.get(t)
        if not k or k["exp"] < time.time(): return None
        for u in users():
            if u["u"] == k["u"]: return u
        return None
    def ip(s):
        if os.environ.get("TRUST_PROXY") == "1": return (s.headers.get("Fly-Client-IP") or (s.headers.get("X-Forwarded-For") or "").split(",")[0].strip() or s.client_address[0])
        return s.client_address[0]
    def q(s): return parse_qs(urlparse(s.path).query)
    def proj(s): return code_of((s.q().get("p") or ["LR001"])[0])
    def site_file(s, path):
        """Public landing-page files (images, videos) with Range support so videos can seek."""
        name = os.path.basename(path); fp = os.path.join(APP, "site", name)
        if not name or name.startswith(".") or not os.path.isfile(fp): s.out(404, ""); return
        ext = name.rsplit(".", 1)[-1].lower()
        ct = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp", "mp4": "video/mp4", "webm": "video/webm", "svg": "image/svg+xml", "css": "text/css", "js": "text/javascript", "ico": "image/x-icon"}.get(ext)
        if not ct: s.out(404, ""); return
        size = os.path.getsize(fp); rng = s.headers.get("Range", ""); a, b, code = 0, size - 1, 200
        m = re.match(r"bytes=(\d*)-(\d*)$", rng)
        if m and (m.group(1) or m.group(2)):
            if m.group(1): a = int(m.group(1)); b = int(m.group(2)) if m.group(2) else size - 1
            else: a = max(0, size - int(m.group(2))); b = size - 1
            b = min(b, size - 1)
            if a > b: s.send_response(416); s.send_header("Content-Range", "bytes */%d" % size); s.send_header("Content-Length", "0"); s.end_headers(); return
            code = 206
        with open(fp, "rb") as f: f.seek(a); body = f.read(b - a + 1)
        s.send_response(code); s.send_header("Content-Type", ct); s.send_header("Accept-Ranges", "bytes"); s.send_header("Cache-Control", "public, max-age=3600"); s.send_header("X-Content-Type-Options", "nosniff")
        if code == 206: s.send_header("Content-Range", "bytes %d-%d/%d" % (a, b, size))
        s.send_header("Content-Length", str(len(body))); s.end_headers()
        if s.command != "HEAD": s.wfile.write(body)
    def do_GET(s):
        path = urlparse(s.path).path
        if path.startswith("/photos/"):
            fp = os.path.join(BASE, "photos", os.path.basename(path))
            if os.path.exists(fp): s.out(200, open(fp, "rb").read(), "image/jpeg")
            else: s.out(404, "")
        elif path.startswith("/logo/"):
            fp = os.path.join(BASE, os.path.basename(path))
            ext = fp.rsplit(".", 1)[-1]
            if os.path.basename(path).startswith("logo.") and os.path.exists(fp): s.out(200, open(fp, "rb").read(), {"png": "image/png", "jpg": "image/jpeg", "webp": "image/webp"}.get(ext, "application/octet-stream"))
            else: s.out(404, "")
        elif path == "/healthz": s.out(200, '{"ok":true}')
        elif path == "/api/public":
            c = company(); s.out(200, json.dumps({"name": c.get("name", ""), "logo": c.get("logo", ""), "copyright": COPYRIGHT_HOLDER}))
        elif path.startswith("/api/"):
            u = s.me()
            if path == "/api/ping": s.out(200, '{"ok":true}'); return
            if not u: s.out(401, '{"error":"login required"}'); return
            if path == "/api/me": s.out(200, json.dumps(pub(u)))
            elif path == "/api/projects": s.out(200, json.dumps(list_projects()))
            elif path == "/api/company": s.out(200, json.dumps(company()))
            elif path == "/api/status":
                s.out(200, json.dumps({"email": bool(os.environ.get("SMTP_HOST")), "whatsapp": wa_provider(), "ai": bool(os.environ.get("ANTHROPIC_API_KEY")), "chan": company().get("chan") or "email", "backups": len(glob.glob(os.path.join(BASE, "backups", "*.zip"))), "demo": os.environ.get("SEED_DEMO_USERS", "1") != "0", "persistent": is_persistent()}))
            elif path == "/api/project/archive":
                if u["role"] not in ("Management", "Project manager"): s.out(403, '{"error":"Management or the project manager only"}'); return
                c_ = code_of((s.q().get("p") or [""])[0])
                if not project_exists(c_): s.out(404, '{"error":"unknown project"}'); return
                why = (s.q().get("why") or ["archive"])[0][:80]
                with LOCK: fn, data = save_archive(c_, "archives", u["u"], why)
                audit("project archive %s (%s)" % (fn, why), u["u"], s.ip())
                if (s.q().get("save") or [""])[0] == "1": s.out(200, json.dumps({"file": fn, "bytes": len(data)})); return
                s.out(200, data, "application/zip", {"Content-Disposition": 'attachment; filename="%s"' % fn}); return
            elif path == "/api/backup/status":
                if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
                if (s.q().get("run") or [""])[0] == "1": daily_backup(True); audit("backup now", u["u"], s.ip())
                s.out(200, json.dumps(backup_status())); return
            elif path == "/api/backup":
                if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
                audit("backup download", u["u"], s.ip()); s.out(200, make_backup(s.q().get("photos") == ["1"]), "application/zip")
            elif path == "/api/users":
                if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
                s.out(200, json.dumps([pub(x) for x in users()]))
            elif path == "/api/state":
                with LOCK:
                    try:
                        p = s.proj()
                        if not project_exists(p):
                            if p == "LR001" and os.path.exists(os.path.join(APP, "seed_state.json")): write_state(json.load(open(os.path.join(APP, "seed_state.json"))), p)
                            else: s.out(404, '{"error":"unknown project"}'); return
                        s.out(200, json.dumps(read_state(p) if u["role"] in PRICE_ROLES else strip_prices(read_state(p))))
                    except sqlite3.OperationalError: s.out(423, '{"error":"database busy, try again"}')
            elif path.startswith("/api/docs/"): s.docs_get(path, u)
            elif path == "/api/engrec":
                if u["role"] == "Main contractor": s.out(403, '{"error":"not allowed"}'); return
                p = s.proj()
                if (s.q().get("verify") or [""])[0]: s.out(200, json.dumps(engrec_verify(p))); return
                s.out(200, json.dumps(engrec_all(p))); return
            elif path == "/api/audit":
                if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
                s.out(200, json.dumps(audit_tail(int((s.q().get("n") or ["500"])[0] or 500)))); return
            elif path == "/api/lib/list":
                s.out(200, json.dumps({"cats": lib.CATS, "items": lib.listing(s.proj(), u["role"] == "Main contractor")})); return
            elif path == "/api/lib/file":
                fl = lib.get(s.proj(), int((s.q().get("id") or ["0"])[0] or 0), u["role"] == "Main contractor")
                if not fl: s.out(404, '{"error":"not found"}'); return
                s.out(200, fl[2], fl[1], {"Content-Disposition": 'attachment; filename="%s"' % re.sub(r'[^A-Za-z0-9._-]', "_", fl[0])}); return
            elif path == "/api/app":
                s.out(200, app_payload(u["role"]), "text/html; charset=utf-8", {"Cache-Control": "no-store"}); return
            elif path == "/api/changelog":
                fp = os.path.join(APP, "legal", "release_history.json")
                s.out(200, open(fp, "rb").read() if os.path.exists(fp) else "[]"); return
            elif path == "/api/legal":
                nm = {"LICENSE": "LICENSE.txt", "COPYRIGHT": "COPYRIGHT.txt", "THIRD_PARTY": "THIRD_PARTY_NOTICES.txt"}.get((s.q().get("f") or [""])[0])
                fp = os.path.join(APP, nm or "-")
                if not nm or not os.path.exists(fp): s.out(404, '{"error":"not found"}'); return
                s.out(200, open(fp, "rb").read(), "text/plain; charset=utf-8"); return
            elif path == "/api/manual":
                fp = os.path.join(APP, "FenOps_User_Manual.pdf")
                if not os.path.exists(fp): s.out(404, '{"error":"not found"}'); return
                s.out(200, open(fp, "rb").read(), "application/pdf", {"Content-Disposition": 'attachment; filename="FenOps_User_Manual.pdf"'}); return
            elif path == "/api/watch/status":
                if u["role"] not in ("Management", "Project manager", "Planning engineer"): s.out(403, '{"error":"not allowed"}'); return
                s.out(200, json.dumps(watch.status(DB_FILE, s.proj(), read_state(s.proj())))); return
            else: s.out(404, "{}")
        elif path in ("/welcome", "/welcome/"): s.out(200, open(os.path.join(APP, "site", "index.html"), "rb").read(), "text/html; charset=utf-8", {"Cache-Control": "no-cache"})
        elif path.startswith("/site/"): s.site_file(path)
        elif path in ASSETS: f, ct = ASSETS[path]; s.out(200, open(os.path.join(APP, f), "rb").read(), ct)
        elif os.environ.get("APP_DEV") == "1" and os.path.exists(os.path.join(APP, "index.html")): s.out(200, open(os.path.join(APP, "index.html"), "rb").read(), "text/html; charset=utf-8")
        else: s.out(200, open(os.path.join(APP, "shell.html"), "rb").read(), "text/html; charset=utf-8", {"Cache-Control": "no-cache"})
    def do_POST(s):
        path = urlparse(s.path).path
        n = int(s.headers.get("Content-Length", 0) or 0)
        if n > 12_000_000: s.out(413, '{"error":"too large"}'); return
        B = json.loads(s.rfile.read(n) or b"{}")
        if path == "/api/signup":
            un = str(B.get("u", "")).strip(); pw = str(B.get("password", "")); role = B.get("role", "Site engineer"); em = str(B.get("email", "")).strip()
            if not re.match(r"^[A-Za-z0-9_.-]{3,30}$", un): s.out(400, '{"error":"User name: 3 to 30 letters, numbers, dot, dash or underscore"}'); return
            if len(pw) < 8: s.out(400, '{"error":"Password needs 8 or more characters"}'); return
            if role not in KNOWN_ROLES or role == "Management": s.out(400, '{"error":"Choose a role"}'); return
            if em and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", em): s.out(400, '{"error":"Email address looks wrong"}'); return
            now = time.time(); SIGNUPS[:] = [t for t in SIGNUPS if now - t < 600]
            if len(SIGNUPS) >= 10: s.out(429, '{"error":"Too many sign-ups, try again later"}'); return
            with LOCK:
                us = users()
                if any(x["u"].lower() == un.lower() for x in us): s.out(400, '{"error":"That user name is taken"}'); return
                if sum(1 for x in us if x.get("status") == "pending") >= 50: s.out(429, '{"error":"Too many requests waiting for approval"}'); return
                sa, h = hash_pw(pw); us.append({"u": un, "name": str(B.get("name", un))[:80] or un, "role": role, "email": em[:120], "phone": str(B.get("phone", ""))[:30], "salt": sa, "hash": h, "must": False, "status": "pending"}); save_json(USERS_FILE, us); SIGNUPS.append(now)
            s.out(200, '{"ok":true}'); return
        if path == "/api/login":
            ip = s.ip(); un = str(B.get("u", "")).strip().lower(); lk = locked("u:" + un) or locked("i:" + ip, 40)
            if lk: audit("login blocked", un, ip); s.out(429, json.dumps({"error": "Too many failed attempts. Try again in %d minute(s)." % (lk // 60 + 1)})); return
            for u in users():
                if u["u"].lower() == un:
                    if hash_pw(str(B.get("p", "")), u["salt"])[1] == u["hash"]:
                        if u.get("status") == "pending": s.out(403, '{"error":"Your account is waiting for approval by Management."}'); return
                        t = secrets.token_hex(24); TOKENS[t] = {"u": u["u"], "exp": time.time() + 12 * 3600}; FAILS.pop("u:" + un, None); audit("login", u["u"], ip)
                        s.out(200, json.dumps({"token": t, "user": pub(u)})); return
            FAILS.setdefault("u:" + un, []).append(time.time()); FAILS.setdefault("i:" + ip, []).append(time.time()); audit("login failed", un, ip)
            time.sleep(0.6); s.out(401, '{"error":"wrong user name or password"}'); return
        u = s.me()
        if not u: s.out(401, '{"error":"login required"}'); return
        if path == "/api/passwd":
            if hash_pw(str(B.get("old", "")), u["salt"])[1] != u["hash"]: s.out(400, '{"error":"current password is wrong"}'); return
            if len(str(B.get("new", ""))) < 8: s.out(400, '{"error":"use at least 8 characters"}'); return
            with LOCK:
                us = users()
                for x in us:
                    if x["u"] == u["u"]: x["salt"], x["hash"] = hash_pw(B["new"]); x["must"] = False
                save_json(USERS_FILE, us)
            s.out(200, "{}"); return
        if path == "/api/users":
            audit("users " + str(B.get("op", "list")), u["u"] + " -> " + str(B.get("u", "")), s.ip())
            if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
            with LOCK:
                us = users(); op = B.get("op"); un = str(B.get("u", "")).strip()
                if op == "add":
                    if not un or any(x["u"].lower() == un.lower() for x in us): s.out(400, '{"error":"user name missing or already used"}'); return
                    if len(str(B.get("password", ""))) < 8: s.out(400, '{"error":"password needs 8+ characters"}'); return
                    sa, h = hash_pw(B["password"]); us.append({"u": un, "name": B.get("name", un), "role": B.get("role", "Site engineer"), "email": B.get("email", ""), "phone": B.get("phone", ""), "salt": sa, "hash": h, "must": True})
                else:
                    x = next((y for y in us if y["u"] == un), None)
                    if not x: s.out(404, '{"error":"no such user"}'); return
                    mgrs = [y for y in us if y["role"] == "Management"]
                    if x.get("super") and op in ("del", "reject"): s.out(400, '{"error":"the super admin cannot be removed"}'); return
                    if x.get("super") and op == "set" and B.get("role", x["role"]) != "Management": s.out(400, '{"error":"the super admin stays Management"}'); return
                    if op == "del":
                        if x["role"] == "Management" and len(mgrs) < 2: s.out(400, '{"error":"keep at least one Management user"}'); return
                        us.remove(x)
                    elif op == "approve":
                        x["status"] = "active"
                        if B.get("role") in KNOWN_ROLES: x["role"] = B["role"]
                    elif op == "reject":
                        if x.get("status") != "pending": s.out(400, '{"error":"only pending requests can be rejected"}'); return
                        us.remove(x)
                    elif op == "reset":
                        if len(str(B.get("password", ""))) < 8: s.out(400, '{"error":"password needs 8+ characters"}'); return
                        x["salt"], x["hash"] = hash_pw(B["password"]); x["must"] = True
                    elif op == "set":
                        if x["role"] == "Management" and B.get("role", x["role"]) != "Management" and len(mgrs) < 2: s.out(400, '{"error":"keep at least one Management user"}'); return
                        for k in ("name", "role", "email", "phone"):
                            if k in B: x[k] = B[k]
                save_json(USERS_FILE, us)
            s.out(200, json.dumps([pub(x) for x in users()])); return
        if path == "/api/photo":
            d = B.get("data", ""); raw = base64.b64decode(d.split(",", 1)[1]) if "," in d else b""
            os.makedirs(os.path.join(BASE, "photos"), exist_ok=True); fn = uuid.uuid4().hex[:12] + ".jpg"
            open(os.path.join(BASE, "photos", fn), "wb").write(raw); s.out(200, json.dumps({"url": "/photos/" + fn})); return
        if path == "/api/company":
            if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
            with LOCK:
                c = company()
                for k in CO_FIELDS:
                    if k in B: c[k] = str(B[k])[:300]
                if c.get("chan") not in (None, "email", "whatsapp", "both"): c["chan"] = "email"
                if isinstance(B.get("perms"), dict):
                    c["perms"] = {r: [t for t in v if isinstance(t, str) and re.match(r"^[a-z0-9]{2,12}$", t)] for r, v in B["perms"].items() if r in KNOWN_ROLES and r != "Management" and isinstance(v, list)}
                elif B.get("perms") is None and "perms" in B: c.pop("perms", None)
                save_json(COMPANY_FILE, c)
            s.out(200, json.dumps(company())); return
        if path == "/api/logo":
            if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
            d = B.get("data", "")
            with LOCK:
                c = company()
                for f in glob.glob(os.path.join(BASE, "logo.*")): os.remove(f)
                if d:
                    m = re.match(r"^data:(image/(?:png|jpeg|webp));base64,", d)
                    if not m: s.out(400, '{"error":"use a PNG, JPG or WEBP image"}'); return
                    fn = "logo." + LOGO_EXT[m.group(1)]; open(os.path.join(BASE, fn), "wb").write(base64.b64decode(d.split(",", 1)[1])); c["logo"] = "/logo/%s?v=%d" % (fn, int(time.time()))
                else: c["logo"] = ""
                save_json(COMPANY_FILE, c)
            s.out(200, json.dumps({"logo": c["logo"]})); return
        if path == "/api/xer":
            try:
                import xer_detail
                if not isinstance(B.get("acts"), list) or not B["acts"] or len(B["acts"]) > 5000: s.out(400, '{"error":"No activities to export"}'); return
                data = xer_detail.build(B)
            except Exception as ex: s.out(400, json.dumps({"error": "Could not build the XER: %s" % ex})); return
            s.out(200, data, "text/plain; charset=windows-1252"); return
        if path == "/api/ask":
            key = os.environ.get("ANTHROPIC_API_KEY")
            if not key: s.out(501, '{"error":"AI guide is not configured on this server"}'); return
            now = time.time(); ASKS[u["u"]] = [t for t in ASKS.get(u["u"], []) if now - t < 60]
            if len(ASKS[u["u"]]) >= 12: s.out(429, '{"error":"Slow down: too many questions this minute"}'); return
            ASKS[u["u"]].append(now)
            q = str(B.get("q", ""))[:600]; ctx = str(B.get("ctx", ""))[:9000]
            if not q.strip(): s.out(400, '{"error":"empty question"}'); return
            import urllib.request as ur
            body = json.dumps({"model": os.environ.get("ANTHROPIC_MODEL", "claude-sonnet-5-5"), "max_tokens": 700, "system": "You are the in-app guide for FenOps, a fenestration installation planning and monitoring app. Answer only from the guide text below. Be concise, use short numbered steps, name the tab to open, and say so if the guide does not cover the question. Never invent features.\n\nGUIDE:\n" + ctx, "messages": [{"role": "user", "content": q}]}).encode()
            try:
                rq = ur.Request("https://api.anthropic.com/v1/messages", body, {"content-type": "application/json", "x-api-key": key, "anthropic-version": "2023-06-01"})
                a = json.loads(ur.urlopen(rq, timeout=30).read()); txt = "".join(c.get("text", "") for c in a.get("content", []))
                s.out(200, json.dumps({"answer": txt}))
            except Exception as x: s.out(502, json.dumps({"error": "AI service unavailable"}))
            return
        if path == "/api/ai/contra":
            # v25: AI check of the project handover package for contradictions (IFC, specifications, BOQ, schedule)
            if u["role"] not in ("Management", "Project manager", "Engineering manager", "Planning engineer"): s.out(403, '{"error":"not allowed"}'); return
            key = os.environ.get("ANTHROPIC_API_KEY")
            if not key: s.out(501, '{"error":"AI is not configured on this server (set ANTHROPIC_API_KEY)"}'); return
            now = time.time(); ASKS[u["u"]] = [t for t in ASKS.get(u["u"], []) if now - t < 60]
            if len(ASKS[u["u"]]) >= 6: s.out(429, '{"error":"Slow down: too many AI checks this minute"}'); return
            ASKS[u["u"]].append(now)
            content = []; total = 0; used = []; info = []
            for i in (B.get("lib") if isinstance(B.get("lib"), list) else [])[:6]:
                try: fl = lib.get(s.proj(), int(i or 0), False)
                except (TypeError, ValueError): fl = None
                if not fl: continue
                fn, mime, raw = fl; total += len(raw)
                if total > 20_000_000: break
                low = fn.lower()
                if mime == "application/pdf" or low.endswith(".pdf"):
                    content.append({"type": "text", "text": "DOCUMENT (file %s):" % fn}); content.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(raw).decode()}}); used.append(fn); info.append({"name": fn, "pages": len(re.findall(rb"/Type\s*/Page[^s]", raw))})
                elif mime.startswith("text/") or low.endswith((".txt", ".csv", ".md", ".json")):
                    content.append({"type": "text", "text": "DOCUMENT (file %s):\n%s" % (fn, raw.decode("utf-8", "replace")[:60000])}); used.append(fn)
                elif mime in ("image/png", "image/jpeg", "image/webp", "image/gif"):
                    content.append({"type": "text", "text": "DOCUMENT (image %s):" % fn}); content.append({"type": "image", "source": {"type": "base64", "media_type": mime, "data": base64.b64encode(raw).decode()}}); used.append(fn)
            content.append({"type": "text", "text": "PROJECT DATA FROM THE APP (BOQ and installation schedule):\n" + str(B.get("ctx", ""))[:60000] + "\n\nFind contradictions BETWEEN the documents of the package (IFC drawings, specifications, BOQ, schedule, manhour budget): different quantities, units, materials, sizes, finishes, standards or dates for the same item. Do not report style issues or missing information, only conflicts you can point to in two named places. Reply with JSON only: {\"items\":[{\"docs\":\"which documents conflict, with sheet or clause\",\"issue\":\"what exactly conflicts\"}]}. If there is none, reply {\"items\":[]}."})
            import urllib.request as ur
            body = json.dumps({"model": os.environ.get("ANTHROPIC_MODEL", "claude-sonnet-5-5"), "max_tokens": 4000, "system": "You are a careful document-control engineer for a facade and fenestration contractor. You compare tender and design documents and report only real contradictions, each tied to the two places that disagree. Never invent content.", "messages": [{"role": "user", "content": content}]}).encode()
            try:
                rq = ur.Request(os.environ.get("ANTHROPIC_API_URL", "https://api.anthropic.com/v1/messages"), body, {"content-type": "application/json", "x-api-key": key, "anthropic-version": "2023-06-01"})
                a = json.loads(ur.urlopen(rq, timeout=150).read()); txt = "".join(c.get("text", "") for c in a.get("content", []))
                i0, i1 = txt.find("{"), txt.rfind("}")
                items = json.loads(txt[i0:i1 + 1]).get("items", []) if i0 >= 0 and i1 > i0 else []
                out = [{"docs": str(x.get("docs", ""))[:300], "txt": str(x.get("issue", x.get("txt", "")))[:900]} for x in items if isinstance(x, dict) and (x.get("issue") or x.get("txt"))][:30]
                audit("ai contra %d items, %d files" % (len(out), len(used)), u["u"], s.ip())
                s.out(200, json.dumps({"items": out, "files": used, "info": info}))
            except Exception as x: s.out(502, json.dumps({"error": "AI service unavailable"}))
            return
        if path == "/api/send":
            now = time.time(); SENDS[u["u"]] = [t for t in SENDS.get(u["u"], []) if now - t < 60]
            if len(SENDS[u["u"]]) >= 40: s.out(429, '{"error":"Too many messages this minute. Wait a moment."}'); return
            SENDS[u["u"]].append(now)
            ch = B.get("channel") if B.get("channel") in ("email", "whatsapp", "both") else None
            r = deliver(str(B.get("to", "")).strip(), str(B.get("phone", "")).strip(), str(B.get("subject", ""))[:200], str(B.get("body", ""))[:8000], ch)
            ok = any(v == "sent" for v in r.values()); audit("send " + json.dumps(r), u["u"], s.ip())
            if ok: s.out(200, json.dumps({"ok": True, "result": r}))
            else: s.out(501, json.dumps({"error": "; ".join("%s %s" % (k, v) for k, v in r.items()) or "Nothing was sent", "result": r}))
            return
        if path == "/api/test":
            if u["role"] != "Management": s.out(403, '{"error":"management only"}'); return
            r = deliver(str(B.get("to", "")).strip(), str(B.get("phone", "")).strip(), "FenOps test message", "This is a test message from FenOps. Email and WhatsApp are working.", B.get("channel") if B.get("channel") in ("email", "whatsapp", "both") else "both")
            s.out(200, json.dumps({"result": r})); return
        if path == "/api/project":
            if u["role"] not in ("Management", "Project manager", "Planning engineer"): s.out(403, '{"error":"not allowed"}'); return
            c = code_of(B.get("code", ""))
            with LOCK:
                if project_exists(c): s.out(400, '{"error":"that project code exists"}'); return
                _tomb(c, False)
                st = B.get("state") or {}; st.setdefault("proj", {"code": c, "name": B.get("name", c)}); write_state(st, c)
                par = code_of(B.get("parent", "")) if B.get("parent") else ""
                if par and project_exists(par): set_tree(c, par, B.get("kind") if B.get("kind") in KINDS else "Sub-project")
            s.out(200, json.dumps(list_projects())); return
        if path == "/api/project/delete":
            if u["role"] != "Management": s.out(403, '{"error":"Only Management can delete a project"}'); return
            c = code_of(B.get("code", ""))
            with LOCK:
                if not project_exists(c): s.out(404, '{"error":"unknown project"}'); return
                if B.get("confirm") != c: s.out(400, '{"error":"type the project code to confirm"}'); return
                tr = get_tree()
                if any((v or {}).get("parent") == c for k, v in tr.items()): s.out(400, '{"error":"This project has sub-projects. Delete or move them first."}'); return
                if len(list_projects()) <= 1: s.out(400, '{"error":"The last project cannot be deleted. Create another one first."}'); return
                try: fn = delete_project(c, u["u"])
                except (OSError, sqlite3.Error) as x: s.out(500, json.dumps({"error": "Could not delete: %s" % x})); return
            audit("project deleted %s (archive %s)" % (c, fn), u["u"], s.ip())
            s.out(200, json.dumps({"projects": list_projects(), "archive": fn})); return
        if path == "/api/ptree":
            if u["role"] not in ("Management", "Project manager", "Planning engineer"): s.out(403, '{"error":"not allowed"}'); return
            c = code_of(B.get("code", "")); par = code_of(B.get("parent", "")) if B.get("parent") else ""
            with LOCK:
                if not project_exists(c): s.out(404, '{"error":"unknown project"}'); return
                if par and not project_exists(par): s.out(400, '{"error":"unknown parent"}'); return
                tr = get_tree(); a = par
                while a:
                    if a == c: s.out(400, '{"error":"a project cannot sit under itself or its own sub-projects"}'); return
                    a = (tr.get(a) or {}).get("parent", "")
                set_tree(c, par, B.get("kind") if B.get("kind") in KINDS else (tr.get(c) or {}).get("kind", "Project"))
            s.out(200, json.dumps(list_projects())); return
        if tombstoned(s.proj()) and not project_exists(s.proj()): s.out(410, '{"error":"This project was deleted"}'); return
        if u["role"] == "Main contractor": s.out(403, '{"error":"read only"}'); return
        if path.startswith("/api/docs/"): s.docs_post(path, u, B); return
        if path in ("/api/lib/upload", "/api/lib/delete"):
            if path == "/api/lib/upload":
                if u["role"] == "Main contractor": s.out(403, '{"error":"not allowed"}'); return
                try:
                    d = str(B.get("data", "")); raw = base64.b64decode(d.split(",", 1)[1] if "," in d else d)
                    i = lib.add(s.proj(), B.get("category"), B.get("title"), B.get("filename"), B.get("mime"), raw, B.get("rev"), bool(B.get("shared")) and u["role"] in ("Management", "Project manager"), u["u"])
                except (ValueError, KeyError) as e: s.out(400, json.dumps({"error": str(e)})); return
                audit("library upload %s %s" % (i, B.get("filename")), u["u"], s.ip()); s.out(200, json.dumps({"id": i})); return
            i = int(B.get("id", 0) or 0)
            if u["role"] not in ("Management", "Project manager") and lib.owner(s.proj(), i) != u["u"]: s.out(403, '{"error":"not allowed"}'); return
            lib.remove(s.proj(), i); audit("library delete %s" % i, u["u"], s.ip()); s.out(200, "{}"); return
        p = s.proj()
        if path == "/api/engrec":
            try: r = engrec_add(p, B, u["u"], u["role"]); s.out(200, json.dumps({"n": r["n"]}))
            except OSError as x: s.out(500, json.dumps({"error": str(x)}))
            return
        if path == "/api/chat":
            with LOCK:
                try:
                    cur = read_state(p); m = {k: B.get(k, "") for k in ("d", "tm", "proj", "by", "text")}
                    m["id"] = max([c.get("id", 0) for c in cur["chat"] if isinstance(c.get("id"), int)] + [0]) + 1
                    cur["chat"].append(m); write_state(cur, p); s.out(200, json.dumps(cur["chat"]))
                except sqlite3.OperationalError: s.out(423, '{"error":"database busy, try again"}')
            return
        if path == "/api/state":
            if tombstoned(p) and not project_exists(p): s.out(410, '{"error":"This project was deleted"}'); return
            with LOCK:
                try:
                    cur = read_state(p) if project_exists(p) else {"mail": [], "chat": []}
                    old = len(cur["mail"]); B["chat"] = cur["chat"]
                    if u["role"] not in PRICE_ROLES: keep_prices(B, cur)
                    write_state(B, p); notify(B.get("mail", [])[old:]); s.out(200, "{}")
                except sqlite3.OperationalError: s.out(423, '{"error":"database busy, try again"}')
            return
        s.out(404, "{}")

    # ---------------- document engine (numbering, register, formats, reports) ----------------
    def docs_hide(s, u): return () if u["role"] in ("Management", "Project manager") else tuple(docs.COMMERCIAL)
    def docs_file_out(s, fl, inline=False):
        if not fl: s.out(404, '{"error":"no file for this item"}'); return
        fn, mime, blob = fl
        s.out(200, blob, mime, {"Content-Disposition": 'attachment; filename="%s"' % re.sub(r'[^A-Za-z0-9._-]', "_", fn)})
    def docs_get(s, path, u):
        q = lambda k, d="": (s.q().get(k) or [d])[0]; p = s.proj(); hide = s.docs_hide(u)
        try:
            if path == "/api/docs/config":
                have = {m: docs._has(m) for m in ("openpyxl", "pypdf", "reportlab")}; have["libreoffice"] = bool(shutil.which("soffice") or shutil.which("libreoffice"))
                nc = [x for x in docs.numbering_config(p) if x["type"] not in hide]; tp = [t for t in docs.list_templates(p) if t["type"] not in hide]
                s.out(200, json.dumps({"numbering": nc, "templates": tp, "have": have, "can_edit": u["role"] in ("Management", "Project manager")})); return
            if path == "/api/docs/list":
                s.out(200, json.dumps(docs.list_docs(p, q("type"), q("q"), q("status"), q("from"), q("to"), hide=hide))); return
            if path == "/api/docs/sigs":
                s.out(200, json.dumps({k: v for k, v in docs.sigs(p).items() if k.split("|")[0] not in hide})); return
            if path == "/api/docs/get":
                d = docs.get_doc(p, int(q("id", "0")))
                if not d or d["type"] in hide: s.out(404, '{"error":"not found"}'); return
                s.out(200, json.dumps(d)); return
            if path == "/api/docs/file":
                did = int(q("id", "0")); t = docs.doc_type_of(p, did)
                if not t or t in hide: s.out(404, '{"error":"not found"}'); return
                if q("regen") == "1" and u["role"] not in ("Management", "Project manager", "QA/QC"): s.out(403, '{"error":"not allowed"}'); return
                s.docs_file_out(docs.get_file(p, did, q("regen") == "1", q("as"))); return
            if path == "/api/docs/report":
                fmt = q("fmt", "xlsx"); t = q("type")
                if t and t in hide: s.out(403, '{"error":"not allowed"}'); return
                rows = docs.list_docs(p, t, q("q"), q("status"), q("from"), q("to"), hide=hide)
                rows = [r for r in rows if r["status"] != "Reserved" or q("reserved") == "1"]
                title = (docs.DOC_TYPES[t]["label"] + " register") if t in docs.DOC_TYPES else "Document register"
                audit("docs report %s %s %d" % (fmt, t or "all", len(rows)), u["u"], s.ip())
                s.docs_file_out(docs.report(p, fmt, rows, title, q("copies", "orig"))); return
            if path in ("/api/docs/template/file", "/api/docs/template/test", "/api/docs/template/starter"):
                if u["role"] not in ("Management", "Project manager"): s.out(403, '{"error":"not allowed"}'); return
                if path.endswith("starter"): s.docs_file_out(docs.starter(q("type"), q("kind"))); return
                r = docs.template_row(int(q("id", "0")))
                if not r or r["project"] not in ("", p): s.out(404, '{"error":"not found"}'); return
                if path.endswith("file"): s.docs_file_out((r["filename"], "application/octet-stream", bytes(r["blob"]))); return
                fn, mime, blob, unk = docs.test_render(p, r["id"]); s.docs_file_out((fn, mime, blob)); return
        except (ValueError, KeyError) as e: s.out(400, json.dumps({"error": str(e)})); return
        except RuntimeError as e: s.out(501, json.dumps({"error": str(e)})); return
        except Exception as e:
            print("docs error:", repr(e)); s.out(500, json.dumps({"error": "%s: %s" % (type(e).__name__, e)})); return
        s.out(404, "{}")
    def docs_post(s, path, u, B):
        p = s.proj(); admin = u["role"] in ("Management", "Project manager"); who = u.get("name") or u["u"]; hide = s.docs_hide(u)
        try:
            if path == "/api/docs/alloc":
                t = str(B.get("type", ""))
                if t in hide: s.out(403, '{"error":"not allowed"}'); return
                s.out(200, json.dumps(docs.allocate(p, t, B.get("date"), who, B.get("floor") or 0))); return
            if path == "/api/docs/register":
                items = B.get("items") if isinstance(B.get("items"), list) else [{"type": B.get("type"), "doc": B.get("doc") or {}}]
                out = []
                for it in items[:200]:
                    t = str(it.get("type", ""))
                    if t in hide: continue
                    r = docs.register(p, t, it.get("doc") or {}, who); out.append({"type": t, "number": r["number"], "asked": (it.get("doc") or {}).get("number"), "dup": r["dup"]})
                s.out(200, json.dumps({"ok": True, "done": out})); return
            if path == "/api/docs/letter":
                r = docs.allocate(p, "LTR", B.get("date"), who)
                d = {"number": r["number"], "date": B.get("date") or docs.today(), "title": str(B.get("title", ""))[:300], "to": B.get("to", ""), "from": B.get("from") or who, "ref": B.get("ref", ""),
                     "status": "Issued", "body": str(B.get("body", ""))[:20000], "sent": bool(B.get("sent")), "fields": {}, "lists": {}, "sig": "L" + uuid.uuid4().hex[:8], "src": "LTR:" + uuid.uuid4().hex[:10]}
                x = docs.register(p, "LTR", d, who); audit("letter " + r["number"], u["u"], s.ip()); s.out(200, json.dumps({"id": x["id"], "number": x["number"]})); return
            if not admin and path not in ("/api/docs/regen",): s.out(403, '{"error":"Management or the project manager only"}'); return
            if path == "/api/docs/numbering":
                docs.set_numbering(p, str(B.get("type", "")), B.get("pattern"), B.get("start"), B.get("reset"), B.get("next"))
                audit("numbering %s %s start %s" % (B.get("type"), B.get("pattern"), B.get("start")), u["u"], s.ip())
                s.out(200, json.dumps({"numbering": docs.numbering_config(p)})); return
            if path == "/api/docs/void":
                docs.void_doc(p, int(B.get("id", 0)), B.get("reason"), who); audit("void doc %s" % B.get("id"), u["u"], s.ip()); s.out(200, "{}"); return
            if path == "/api/docs/regen":
                if u["role"] not in ("Management", "Project manager", "QA/QC"): s.out(403, '{"error":"not allowed"}'); return
                docs.make_file(p, int(B.get("id", 0))); s.out(200, "{}"); return
            if path == "/api/docs/template":
                d = str(B.get("data", "")); raw = base64.b64decode(d.split(",", 1)[1] if "," in d else d)
                res = docs.add_template(p, B.get("scope") != "project", str(B.get("type", "")), str(B.get("name", "")), str(B.get("filename", "")), raw, B.get("opts") or {}, who)
                audit("template %s %s" % (B.get("type"), B.get("filename")), u["u"], s.ip()); s.out(200, json.dumps(res)); return
            if path == "/api/docs/template/active":
                docs.set_template_active(p, int(B.get("id", 0)), bool(B.get("active"))); s.out(200, "{}"); return
            if path == "/api/docs/template/delete":
                docs.delete_template(p, int(B.get("id", 0))); audit("template deleted %s" % B.get("id"), u["u"], s.ip()); s.out(200, "{}"); return
        except (ValueError, KeyError) as e: s.out(400, json.dumps({"error": str(e)})); return
        except RuntimeError as e: s.out(501, json.dumps({"error": str(e)})); return
        except Exception as e:
            print("docs error:", repr(e)); s.out(500, json.dumps({"error": "%s: %s" % (type(e).__name__, e)})); return
        s.out(404, "{}")
    def log_message(s, *a): pass

def is_persistent():
    """False when DATA_DIR is on the container's own disk (data is lost on every redeploy/restart)."""
    if not os.environ.get("DATA_DIR"): return True
    try: return os.stat(BASE).st_dev != os.stat("/").st_dev or os.environ.get("ASSUME_PERSISTENT") == "1"
    except OSError: return True
def make_backup(photos=False):
    import io, tempfile, zipfile
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for f in glob.glob(os.path.join(BASE, "*")):
            if os.path.isfile(f) and not f.endswith((".tmp", ".bak")) and not os.path.basename(f).startswith("fenestra.db"): z.write(f, os.path.basename(f))
        src_ = connect()   # consistent snapshot via SQLite's backup API (safe while the app is writing)
        try:
            with tempfile.TemporaryDirectory() as td:
                snap = os.path.join(td, "fenestra.db"); dst = sqlite3.connect(snap); src_.backup(dst); dst.close()
                z.write(snap, "fenestra.db")
            z.writestr("fenestra.sql", "\n".join(src_.iterdump()) + "\n")   # plain SQL text, restorable with: sqlite3 new.db < fenestra.sql
        finally: src_.close()
        if photos:
            for f in glob.glob(os.path.join(BASE, "photos", "*")): z.write(f, "photos/" + os.path.basename(f))
    return b.getvalue()
def check_backup(fn):
    # a backup only counts if the zip is readable and the database inside passes SQLite's integrity check
    import zipfile, tempfile
    try:
        with zipfile.ZipFile(fn) as z:
            if z.testzip() is not None or "fenestra.db" not in z.namelist(): return False
            with tempfile.TemporaryDirectory() as td:
                z.extract("fenestra.db", td); c = sqlite3.connect(os.path.join(td, "fenestra.db"))
                try: return c.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
                finally: c.close()
    except Exception: return False
def daily_backup(force=False):
    d = os.path.join(BASE, "backups"); os.makedirs(d, exist_ok=True); fn = os.path.join(d, time.strftime("%Y-%m-%d") + ".zip")
    if os.path.exists(fn) and not force: return
    with LOCK: data = make_backup(False)
    open(fn + ".tmp", "wb").write(data); os.replace(fn + ".tmp", fn)
    if not check_backup(fn): print("BACKUP CHECK FAILED:", fn)
    for old in sorted(glob.glob(os.path.join(d, "*.zip")))[:-30]: os.remove(old)
def backup_status():
    d = os.path.join(BASE, "backups"); L = sorted(glob.glob(os.path.join(d, "*.zip")))
    if not L: return {"last": None, "n": 0, "age_h": None}
    f = L[-1]; return {"last": {"name": os.path.basename(f), "kb": os.path.getsize(f) / 1024, "ok": check_backup(f)}, "n": len(L), "age_h": (time.time() - os.path.getmtime(f)) / 3600}

def watch_loop():
    # every 10 minutes: delegation alarms, daily project-manager reminders and unassigned-task notices (see watch.py)
    while True:
        try:
            if os.environ.get("SMTP_HOST") or wa_provider():   # without email/WhatsApp configured nothing is sent (the app still shows the alerts)
                watch.run_once(DB_FILE, [p["code"] for p in list_projects()], read_state, deliver, users(), project_name=project_name)
        except Exception as x: print("watch error:", repr(x))
        time.sleep(600)

def saturday_loop():
    # every Saturday (server clock), email site engineers who have an address, if SMTP is configured
    while True:
        try:
            now = dt.datetime.now()
            if now.weekday() == 5 and now.hour >= 7 and (os.environ.get("SMTP_HOST") or wa_provider()):
                with LOCK:
                    c = company(); day = now.strftime("%Y-%m-%d")
                    if c.get("pph_last") != day:
                        c["pph_last"] = day; save_json(COMPANY_FILE, c); todo = [x for x in users() if x.get("role") == "Site engineer" and (x.get("email") or x.get("phone")) and x.get("status", "active") == "active"]
                    else: todo = []
                for x in todo: deliver(x.get("email"), x.get("phone"), "Weekly progress photos", "It is Saturday. Please open FenOps > Progress photos and upload this week's photos, tagged by building and elevation.")
            if now.hour >= 2: daily_backup()
        except Exception: pass
        time.sleep(600)

if __name__ == "__main__":
    port = int(sys.argv[sys.argv.index("--port") + 1]) if "--port" in sys.argv else int(os.environ.get("PORT", 8000))
    users(); init_db(); migrate_legacy()
    docs.init(DB_FILE, company, project_name, logo_path)
    if not project_exists("LR001") and os.path.exists(os.path.join(APP, "seed_state.json")): write_state(json.load(open(os.path.join(APP, "seed_state.json"))), "LR001")
    print(f"Open http://localhost:{port}" + ("   LAN: http://<this-laptop-ip>:%d" % port if "--lan" in sys.argv else ""))
    if os.environ.get("SEED_DEMO_USERS", "1") != "0": print("First logins: user names management, pm, planner, siteeng, qe, hse, contractor, production. Password azal2026 (you must change it at first login).")
    try:
        if any(x.get("must") and x.get("status", "active") == "active" for x in users()): print("NOTICE: some accounts still have their initial password. Set SEED_DEMO_USERS=0 on a live server and reset or remove unused accounts (Management > Users).")
    except Exception: pass
    threading.Thread(target=lambda: (time.sleep(20), daily_backup()), daemon=True).start()
    print("Database:", DB_FILE)
    print("Email:", "on" if os.environ.get("SMTP_HOST") else "off", "| WhatsApp:", wa_provider() or "off", "| AI guide:", "on" if os.environ.get("ANTHROPIC_API_KEY") else "off")
    if not is_persistent(): print("WARNING: %s is not a mounted volume. Users and project data will be LOST on every redeploy or restart. Attach a persistent volume at %s (set ASSUME_PERSISTENT=1 if it is persistent but not detected)." % (BASE, BASE))
    watch.init(DB_FILE); lib.init(DB_FILE)
    threading.Thread(target=watch_loop, daemon=True).start()
    threading.Thread(target=saturday_loop, daemon=True).start()
    ThreadingHTTPServer.daemon_threads = True
    ThreadingHTTPServer(("0.0.0.0" if ("--lan" in sys.argv or os.environ.get("PORT")) else "127.0.0.1", port), H).serve_forever()
